"""Profiles perpendicular to a boundary drawn on an image, cleaned and averaged.

This follows ``greyvalues.m`` of NIDIS (Petrone et al. 2016, code at
https://github.com/cpetrone/NIDIS): the user draws a guideline along the zone
boundary, and values are read along many parallel lines perpendicular to it,
one line per pixel of guideline, then averaged position by position. As in
NIDIS the first value of every line lies to the **right of the guideline,
looking from its first point to its last**, and the guideline sits where the
two lengths meet. NIDIS reads 50 px either side by default and so does this.

What is added to NIDIS:

* the guideline can be a polyline, so a curved zone boundary can be followed;
  each line is perpendicular to the local direction of the guideline;
* any image :mod:`.images` reads, any value mode (grey, one channel, a colour
  legend), and bilinear sampling (MATLAB's ``improfile`` default is nearest);
* cleaning in two stages. In the image: values below or above chosen limits
  (cracks and holes are dark, oxide or sulphide inclusions bright in BSE),
  colours not on the legend, and polygons drawn around inclusions or lamellae,
  all optionally grown by a few pixels because the edge of a crack is a blend of
  crack and crystal. Across the lines: NIDIS's mean +/- 1 SD test at each
  position (it removes about a third of perfectly good Gaussian values, so it is
  offered as a preset, not the default) or a repeated median +/- k MAD test;
  and lines with too many rejected values dropped whole, since a lamella lying
  along a line spoils the whole line;
* every rejected value keeps a reason code, so the workbook shows why.

The result is written to an Excel workbook (raw and cleaned lines, statistics,
geometry, settings) that Diffusor reads back with :func:`read_extraction` and
converts to composition with :func:`composition_table` once a pixel size and a
grey-to-composition map are given.
"""
from __future__ import annotations

import json
import math
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from .greyscale import GreyscaleCalibration, apply_calibration
from .images import ColorScale, LoadedImage, VALUE_MODES, value_map

MARKER = "diffusor image profile"
WORKBOOK_VERSION = 1

# why a value was not used; stored per value in the workbook
KEPT, OUTSIDE, THRESHOLD, EXCLUDED, OFF_SCALE, OUTLIER, LINE_DROPPED = range(7)
REASONS = {
    KEPT: "kept",
    OUTSIDE: "outside the image",
    THRESHOLD: "below or above the value limits (or next to such a pixel)",
    EXCLUDED: "inside a drawn exclusion area",
    OFF_SCALE: "pixel has no value (colour not on the legend, or NaN)",
    OUTLIER: "outlier across the lines at this position",
    LINE_DROPPED: "line dropped: too many of its values rejected",
}

CLIP_METHODS = {
    "none": "No outlier test",
    "mad": "Median +/- k MAD, repeated",
    "nidis": "Mean +/- k SD, once (NIDIS)",
}


@dataclass
class CleaningSettings:
    low: Optional[float] = None               # values below are impurities (cracks, holes)
    high: Optional[float] = None              # values above are impurities (bright inclusions)
    grow_px: int = 0                          # grow rejected image areas by this radius
    exclusions: List[List[Tuple[float, float]]] = field(default_factory=list)  # polygons (x, y)
    clip_method: str = "mad"                  # see CLIP_METHODS
    clip_k: float = 3.0
    clip_iterations: int = 5
    max_rejected_fraction: Optional[float] = 0.5   # drop a line with more rejected than this

    def describe(self) -> str:
        parts = []
        if self.low is not None or self.high is not None:
            lo = "-inf" if self.low is None else f"{self.low:g}"
            hi = "+inf" if self.high is None else f"{self.high:g}"
            parts.append(f"values outside [{lo}, {hi}] rejected")
        if self.exclusions:
            parts.append(f"{len(self.exclusions)} exclusion area(s)")
        if self.grow_px and (parts or self.exclusions):
            parts.append(f"rejected areas grown by {self.grow_px} px")
        if self.clip_method == "nidis":
            parts.append(f"mean +/- {self.clip_k:g} SD at each position, once (NIDIS)")
        elif self.clip_method == "mad":
            parts.append(f"median +/- {self.clip_k:g} x 1.4826 MAD at each position, "
                         f"up to {self.clip_iterations} passes")
        if self.max_rejected_fraction is not None:
            parts.append(f"lines with more than {100 * self.max_rejected_fraction:.0f}% "
                         "rejected dropped")
        return "; ".join(parts) or "no cleaning"


@dataclass
class ExtractionSettings:
    guideline: List[Tuple[float, float]]      # (x, y) pixel coordinates, x right, y down
    length_before_px: float = 50.0            # on the right of the guideline (NIDIS hp)
    length_after_px: float = 50.0             # on the left
    line_spacing_px: float = 1.0              # distance between lines along the guideline
    sample_step_px: float = 1.0               # distance between values along a line
    interpolation: str = "bilinear"           # or "nearest", as MATLAB improfile
    flip: bool = False                        # start on the left instead
    value_mode: str = "luminance"             # see images.VALUE_MODES
    channel: int = 0
    pixel_size_um: Optional[float] = None
    cleaning: CleaningSettings = field(default_factory=CleaningSettings)

    def to_json(self) -> str:
        return json.dumps(asdict(self))

    @classmethod
    def from_json(cls, text: str) -> "ExtractionSettings":
        d = json.loads(text)
        d["cleaning"] = CleaningSettings(**d.get("cleaning", {}))
        d["guideline"] = [tuple(p) for p in d["guideline"]]
        return cls(**d)


# ------------------------------------------------------------------ geometry
def _polyline(points) -> Tuple[np.ndarray, np.ndarray]:
    pts = np.asarray(points, dtype=float)
    if pts.ndim != 2 or pts.shape[1] != 2 or len(pts) < 2:
        raise ValueError("the guideline needs at least two points")
    keep = np.r_[True, np.hypot(*np.diff(pts, axis=0).T) > 1e-9]
    pts = pts[keep]
    if len(pts) < 2:
        raise ValueError("the guideline points coincide")
    arc = np.r_[0.0, np.cumsum(np.hypot(*np.diff(pts, axis=0).T))]
    return pts, arc


def _at(pts, arc, s) -> np.ndarray:
    return np.column_stack([np.interp(s, arc, pts[:, 0]), np.interp(s, arc, pts[:, 1])])


def stations(points, spacing: float = 1.0, smooth_px: float = 2.0):
    """Positions, unit directions and arc lengths of the lines along the guideline.

    Lines are evenly spaced from the first point to the last. The direction at
    each station is taken over +/- ``smooth_px`` of guideline, so a corner of a
    polyline turns the lines gradually rather than all at once.
    """
    pts, arc = _polyline(points)
    L = float(arc[-1])
    n = max(1, int(round(L / spacing))) + 1
    s = np.linspace(0.0, L, n)
    pos = _at(pts, arc, s)
    h = min(smooth_px, L / 2.0)
    d = _at(pts, arc, np.clip(s + h, 0, L)) - _at(pts, arc, np.clip(s - h, 0, L))
    d /= np.maximum(np.hypot(d[:, 0], d[:, 1]), 1e-12)[:, None]
    return pos, d, s


def line_geometry(settings: ExtractionSettings):
    """Every sample point: arrays (n_samples, n_lines) of x and y, plus the stations.

    Distance along a line runs from its right-hand end (NIDIS) unless ``flip``.
    """
    pos, d, s = stations(settings.guideline, settings.line_spacing_px)
    right = np.column_stack([-d[:, 1], d[:, 0]])          # image y points down
    a, b = float(settings.length_before_px), float(settings.length_after_px)
    if a < 0 or b < 0 or a + b <= 0:
        raise ValueError("the profile lines need a positive length")
    step = float(settings.sample_step_px)
    m = int(round((a + b) / step)) + 1
    dist = np.linspace(0.0, a + b, m)
    offset = dist - a                                      # 0 on the guideline
    # position = station + (a - distance) * right; flipped, the line starts on the left
    along = (dist - a) if settings.flip else (a - dist)
    xs = pos[None, :, 0] + along[:, None] * right[None, :, 0]
    ys = pos[None, :, 1] + along[:, None] * right[None, :, 1]
    return xs, ys, dist, offset, pos, s


# ------------------------------------------------------------------ extraction
@dataclass
class ProfileExtraction:
    settings: ExtractionSettings
    source: str
    image_description: str
    value_label: str
    distance_px: np.ndarray              # (m,) from the start of each line
    offset_px: np.ndarray                # (m,) from the guideline, negative on the start side
    raw: np.ndarray                      # (m, n) values, NaN outside the image
    reason: np.ndarray                   # (m, n) reason codes, KEPT where used
    xs: np.ndarray                       # (m, n) sample x
    ys: np.ndarray                       # (m, n) sample y
    station_xy: np.ndarray               # (n, 2)
    station_arc_px: np.ndarray           # (n,)
    color_scale: Optional[ColorScale] = None
    notes: List[str] = field(default_factory=list)

    @property
    def n_lines(self) -> int:
        return int(self.raw.shape[1])

    @property
    def n_samples(self) -> int:
        return int(self.raw.shape[0])

    @property
    def clean(self) -> np.ndarray:
        return np.where(self.reason == KEPT, self.raw, np.nan)

    @property
    def dropped_lines(self) -> np.ndarray:
        return np.all((self.reason == LINE_DROPPED) | (self.reason == OUTSIDE), axis=0) & \
            np.any(self.reason == LINE_DROPPED, axis=0)

    def distance_um(self) -> Optional[np.ndarray]:
        s = self.settings.pixel_size_um
        return None if not s else self.distance_px * s

    def reason_counts(self) -> Dict[str, int]:
        return {REASONS[k]: int(np.sum(self.reason == k)) for k in REASONS}

    def statistics(self, which: str = "clean") -> pd.DataFrame:
        """NIDIS's per-position statistics of the raw or the cleaned values."""
        v = self.clean if which == "clean" else np.where(self.reason == OUTSIDE, np.nan, self.raw)
        n = np.sum(np.isfinite(v), axis=1)
        with np.errstate(invalid="ignore", divide="ignore"), _quiet():
            mean = np.nanmean(v, axis=1)
            sd = np.nanstd(v, axis=1, ddof=1)
            out = {
                "N": n,
                "Min": np.nanmin(v, axis=1),
                "Max": np.nanmax(v, axis=1),
                "Mean": mean,
                "Median": np.nanmedian(v, axis=1),
                "SD": sd,
                "Rel_SD": sd / np.abs(mean),
                "SE": sd / np.sqrt(n),
            }
        return pd.DataFrame(out)

    def profile_table(self) -> pd.DataFrame:
        """The table Diffusor loads: distances, then raw and clean statistics."""
        cols = {"Distance_px": self.distance_px, "Offset_px": self.offset_px}
        um = self.distance_um()
        if um is not None:
            cols["Distance_um"] = um
            cols["Offset_um"] = self.offset_px * self.settings.pixel_size_um
        df = pd.DataFrame(cols)
        for which, prefix in (("raw", "Raw_"), ("clean", "Clean_")):
            st = self.statistics(which)
            for c in st.columns:
                df[prefix + c] = st[c].to_numpy()
        return df


class _quiet:
    """Silence numpy's all-NaN slice warnings; positions with no data stay NaN."""
    def __enter__(self):
        import warnings
        self._w = warnings.catch_warnings()
        self._w.__enter__()
        warnings.simplefilter("ignore", RuntimeWarning)

    def __exit__(self, *exc):
        self._w.__exit__(*exc)


def _disk(r: int) -> np.ndarray:
    y, x = np.mgrid[-r:r + 1, -r:r + 1]
    return x * x + y * y <= r * r


def impurity_masks(values: np.ndarray, no_value: np.ndarray, cleaning: CleaningSettings):
    """Image-space masks: (below/above limits, inside exclusion polygons), grown."""
    from scipy.ndimage import binary_dilation
    thr = np.zeros(values.shape, dtype=bool)
    with np.errstate(invalid="ignore"):
        if cleaning.low is not None:
            thr |= values < cleaning.low
        if cleaning.high is not None:
            thr |= values > cleaning.high
    exc = np.zeros(values.shape, dtype=bool)
    if cleaning.exclusions:
        from matplotlib.path import Path as MplPath
        h, w = values.shape
        for poly in cleaning.exclusions:
            if len(poly) < 3:
                continue
            pp = np.asarray(poly, float)
            x0, y0 = np.floor(pp.min(axis=0)).astype(int)
            x1, y1 = np.ceil(pp.max(axis=0)).astype(int)
            x0, y0, x1, y1 = max(x0, 0), max(y0, 0), min(x1, w - 1), min(y1, h - 1)
            if x1 < x0 or y1 < y0:
                continue
            yy, xx = np.mgrid[y0:y1 + 1, x0:x1 + 1]
            inside = MplPath(pp).contains_points(np.column_stack([xx.ravel(), yy.ravel()]),
                                                 radius=0.5)
            exc[y0:y1 + 1, x0:x1 + 1] |= inside.reshape(xx.shape)
    off = no_value.copy()
    if cleaning.grow_px and cleaning.grow_px > 0:
        k = _disk(int(cleaning.grow_px))
        if thr.any():
            thr = binary_dilation(thr, k)
        if exc.any():
            exc = binary_dilation(exc, k)
        if off.any():
            off = binary_dilation(off, k)
    return thr, exc, off


def _sample(img: np.ndarray, xs, ys, order: int, cval=np.nan):
    from scipy.ndimage import map_coordinates
    return map_coordinates(img, [ys.ravel(), xs.ravel()], order=order, mode="constant",
                           cval=cval, prefilter=False).reshape(xs.shape)


def _clip(values: np.ndarray, usable: np.ndarray, method: str, k: float, iterations: int):
    """Outliers across the lines at each position, as a boolean array."""
    out = np.zeros(values.shape, dtype=bool)
    if method == "none":
        return out
    v = np.where(usable, values, np.nan)
    passes = 1 if method == "nidis" else max(1, int(iterations))
    for _ in range(passes):
        with _quiet():
            if method == "nidis":
                centre = np.nanmean(v, axis=1)
                spread = np.nanstd(v, axis=1, ddof=1)
            else:
                centre = np.nanmedian(v, axis=1)
                spread = 1.4826 * np.nanmedian(np.abs(v - centre[:, None]), axis=1)
        spread = np.where(np.isfinite(spread), spread, 0.0)
        with np.errstate(invalid="ignore"):
            # a zero spread (flat, saturated or quantised data) rejects nothing
            new = (np.abs(v - centre[:, None]) > k * spread[:, None]) & (spread[:, None] > 0)
        new &= np.isfinite(v)
        if not new.any():
            break
        out |= new
        v = np.where(new, np.nan, v)
    return out


def extract_profiles(image: LoadedImage, settings: ExtractionSettings,
                     color_scale: Optional[ColorScale] = None,
                     values: Optional[Tuple[np.ndarray, np.ndarray, str]] = None
                     ) -> ProfileExtraction:
    """Read, clean and average values along lines perpendicular to the guideline.

    ``values`` may pass a value map already computed with :func:`images.value_map`
    for this image and mode, which saves recomputing a colour-scale inversion
    each time only the lines or the cleaning change.
    """
    if values is None:
        values = value_map(image, settings.value_mode, settings.channel, color_scale)
    vals, no_value, label = values
    cl = settings.cleaning
    xs, ys, dist, offset, pos, arc = line_geometry(settings)
    order = 0 if settings.interpolation == "nearest" else 1
    if order == 0:
        xs_s, ys_s = np.round(xs), np.round(ys)
    else:
        xs_s, ys_s = xs, ys
    filled = np.where(no_value, 0.0, vals)
    raw = _sample(filled, xs_s, ys_s, order)
    h, w = vals.shape
    outside = ~((xs >= -0.5) & (xs <= w - 0.5) & (ys >= -0.5) & (ys <= h - 0.5)) | ~np.isfinite(raw)
    # a sample touching any flagged pixel is flagged: with bilinear sampling the
    # flagged pixel would otherwise leak into the value
    thr, exc, off = impurity_masks(vals, no_value, cl)
    reason = np.full(raw.shape, KEPT, dtype=np.int8)

    def touched(mask):
        return _sample(mask.astype(float), xs_s, ys_s, order, cval=0.0) > 1e-9

    if off.any():
        reason[touched(off)] = OFF_SCALE
    if exc.any():
        reason[(reason == KEPT) & touched(exc)] = EXCLUDED
    if thr.any():
        reason[(reason == KEPT) & touched(thr)] = THRESHOLD
    reason[outside] = OUTSIDE
    raw = np.where(outside | (reason == OFF_SCALE), np.nan, raw)
    outl = _clip(raw, reason == KEPT, cl.clip_method, cl.clip_k, cl.clip_iterations)
    reason[outl & (reason == KEPT)] = OUTLIER
    notes: List[str] = []
    if cl.max_rejected_fraction is not None:
        inside = reason != OUTSIDE
        rejected = inside & (reason != KEPT)
        frac = rejected.sum(axis=0) / np.maximum(inside.sum(axis=0), 1)
        drop = frac > cl.max_rejected_fraction
        if drop.any():
            reason[:, drop] = np.where(reason[:, drop] == OUTSIDE, OUTSIDE, LINE_DROPPED)
            notes.append(f"{int(drop.sum())} of {drop.size} lines dropped: more than "
                         f"{100 * cl.max_rejected_fraction:.0f}% of their values were rejected")
    n_out = int((reason == OUTSIDE).sum())
    if n_out:
        notes.append(f"{n_out} sample points fall outside the image and are not used")
    used = np.sum(reason == KEPT, axis=1)
    if used.size and used.min() < 3:
        notes.append(f"{int(np.sum(used < 3))} positions have fewer than 3 usable values")
    if settings.line_spacing_px < 1.0 or settings.sample_step_px < 1.0:
        notes.append("lines or samples closer than one pixel reuse the same pixels; the "
                     "standard error then understates the real uncertainty")
    ex = ProfileExtraction(settings, image.source, image.describe(), label, dist, offset,
                           raw, reason, xs, ys, pos, arc,
                           color_scale if settings.value_mode == "colour_scale" else None, notes)
    return ex


def suggest_limits(values: np.ndarray, lo_pct: float = 0.5, hi_pct: float = 99.5):
    """Value limits from percentiles, a starting point for the impurity thresholds."""
    v = values[np.isfinite(values)]
    if v.size == 0:
        return None, None
    return float(np.percentile(v, lo_pct)), float(np.percentile(v, hi_pct))


# ------------------------------------------------------------------ workbook
def _methods_text(ex: ProfileExtraction) -> List[str]:
    s = ex.settings
    scale = (f"{s.pixel_size_um:.5g} um per pixel" if s.pixel_size_um
             else "no pixel size (distances in pixels only)")
    lines = [
        f"Values were read from {Path(ex.source).name} ({ex.image_description}) as "
        f"{VALUE_MODES.get(s.value_mode, s.value_mode).lower()}"
        + (f" (channel {s.channel + 1})" if s.value_mode == "channel" else "") + ".",
        f"{ex.n_lines} profile lines were placed every {s.line_spacing_px:g} px along a "
        f"{len(s.guideline)}-point guideline of {ex.station_arc_px[-1]:.1f} px, each "
        f"perpendicular to the local guideline direction and running "
        f"{s.length_before_px:g} px before and {s.length_after_px:g} px after it, sampled "
        f"every {s.sample_step_px:g} px with {s.interpolation} interpolation "
        f"(after greyvalues.m of NIDIS, Petrone et al. 2016).",
        "Distance 0 is the " + ("left" if s.flip else "right") + "-hand end of each line, "
        "looking along the guideline from its first point.",
        f"Scale: {scale}.",
        f"Cleaning: {s.cleaning.describe()}.",
    ]
    if ex.color_scale is not None:
        lines.append(f"Colour scale: {ex.color_scale.source}, colours farther than "
                     f"Delta E {ex.color_scale.max_distance:g} from it rejected.")
    return lines


# workbook colours, matching the result reports
INK, TEAL, GOLD = "243746", "007F82", "D39034"


def _cell(v):
    if isinstance(v, np.generic):
        v = v.item()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _append(ws, values):
    ws.append([_cell(v) for v in values])
    for c in ws[ws.max_row]:
        if isinstance(c.value, str):
            c.data_type = "s"          # file names and notes stay text, even if they start with '='


def write_workbook(path, ex: ProfileExtraction, overlay_png: Optional[bytes] = None) -> str:
    """Write the extraction to .xlsx: profile, raw and clean lines, reasons, geometry.

    The ``Profile`` sheet comes first so that any spreadsheet reader sees the
    averaged profile, and ``Settings`` holds a machine-readable copy of every
    setting so the extraction can be repeated exactly.
    """
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if ex.n_lines + 4 > 16384:
        raise ValueError(f"{ex.n_lines} lines do not fit in an Excel sheet; space them wider")
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=INK)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    def header(ws, row=1):
        for c in ws[row]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = ws.cell(row + 1, 1)

    def frame(ws, df):
        _append(ws, list(df.columns))
        for rec in df.itertuples(index=False, name=None):
            _append(ws, rec)
        header(ws)

    # Profile ----------------------------------------------------------------
    prof = ex.profile_table()
    ws = wb.active
    ws.title = "Profile"
    frame(ws, prof)
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    xcol = list(prof.columns).index("Distance_um" if "Distance_um" in prof else "Distance_px") + 1
    chart = ScatterChart()
    chart.title = "Mean across lines: raw and cleaned"
    chart.x_axis.title = "Distance (µm)" if "Distance_um" in prof else "Distance (px)"
    chart.y_axis.title = ex.value_label
    chart.width, chart.height = 22, 11
    xref = Reference(ws, min_col=xcol, min_row=2, max_row=ws.max_row)
    for name, colour in (("Raw_Mean", GOLD), ("Clean_Mean", TEAL)):
        col = list(prof.columns).index(name) + 1
        ser = Series(Reference(ws, min_col=col, min_row=1, max_row=ws.max_row), xref,
                     title_from_data=True)
        ser.marker.symbol = "circle"
        ser.marker.size = 4
        ser.marker.graphicalProperties.solidFill = colour
        ser.marker.graphicalProperties.line.solidFill = colour
        ser.graphicalProperties.line.noFill = True
        chart.series.append(ser)
    ws.add_chart(chart, get_column_letter(ws.max_column + 2) + "2")

    # Summary ------------------------------------------------------------------
    sm = wb.create_sheet("Summary", 1)
    sm.sheet_properties.tabColor = TEAL
    rows = [("Diffusor image profile",), (),
            ("Source image", ex.source), ("Image", ex.image_description),
            ("Values", ex.value_label), ("Lines", ex.n_lines),
            ("Samples per line", ex.n_samples),
            ("Pixel size (µm/px)", ex.settings.pixel_size_um or "not set"),
            ("Written", datetime.now().isoformat(timespec="seconds")), (),
            ("Methods",)] + [("", t) for t in _methods_text(ex)] + [(), ("Rejected values",)]
    rows += [(k, v) for k, v in ex.reason_counts().items()]
    if ex.notes:
        rows += [(), ("Notes",)] + [("", n) for n in ex.notes]
    rows += [(), ("Using this in Diffusor",),
             ("", "File > Load profile and pick this workbook. Diffusor asks for the pixel size "
                  "(if not set here) and for the map from value to composition: two reference "
                  "points, microprobe anchor points, or none if the values already are "
                  "compositions. It then models the cleaned mean with its standard error.")]
    for r in rows:
        _append(sm, r)
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 100
    for row in sm.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    sm["A1"].font = Font(name="Arial", size=16, bold=True, color=TEAL)
    for row in sm.iter_rows():
        if row[0].value in ("Methods", "Rejected values", "Notes", "Using this in Diffusor"):
            row[0].font = Font(name="Arial", size=11, bold=True, color=INK)
    if overlay_png:
        from io import BytesIO
        from openpyxl.drawing.image import Image as XLImage
        pic = XLImage(BytesIO(overlay_png))
        scale = min(1.0, 640.0 / max(pic.width, 1))
        pic.width, pic.height = pic.width * scale, pic.height * scale
        sm.add_image(pic, "D2")

    # Lines --------------------------------------------------------------------
    names = [f"Line_{i + 1}" for i in range(ex.n_lines)]
    lead = {"Distance_px": ex.distance_px, "Offset_px": ex.offset_px}
    if ex.distance_um() is not None:
        lead["Distance_um"] = ex.distance_um()
    for title, mat in (("Raw lines", np.where(ex.reason == OUTSIDE, np.nan, ex.raw)),
                       ("Clean lines", ex.clean)):
        ws = wb.create_sheet(title)
        df = pd.concat([pd.DataFrame(lead), pd.DataFrame(mat, columns=names)], axis=1)
        frame(ws, df)
    ws = wb.create_sheet("Rejection codes")
    df = pd.concat([pd.DataFrame(lead), pd.DataFrame(ex.reason.astype(int), columns=names)], axis=1)
    frame(ws, df)
    ws = wb.create_sheet("Reason key")
    frame(ws, pd.DataFrame({"Code": list(REASONS), "Meaning": list(REASONS.values()),
                            "Count": [int(np.sum(ex.reason == k)) for k in REASONS]}))
    ws.column_dimensions["B"].width = 62

    # Geometry -----------------------------------------------------------------
    ws = wb.create_sheet("Geometry")
    geo = pd.DataFrame({
        "Line": np.arange(1, ex.n_lines + 1),
        "Arc_along_guideline_px": ex.station_arc_px,
        "Guideline_x": ex.station_xy[:, 0], "Guideline_y": ex.station_xy[:, 1],
        "Start_x": ex.xs[0], "Start_y": ex.ys[0], "End_x": ex.xs[-1], "End_y": ex.ys[-1],
        "Values_used": np.sum(ex.reason == KEPT, axis=0),
        "Dropped": np.where(ex.dropped_lines, "yes", "no"),
    })
    frame(ws, geo)
    r0 = ws.max_row + 2
    _append(ws, ())
    _append(ws, ("Guideline vertex", "x", "y"))
    for i, (x, y) in enumerate(ex.settings.guideline):
        _append(ws, (i + 1, x, y))
    for c in ws[r0]:
        c.font = Font(name="Arial", size=10, bold=True, color=INK)

    if ex.color_scale is not None:
        ws = wb.create_sheet("Colour scale")
        frame(ws, ex.color_scale.table())

    # Settings (machine readable) --------------------------------------------
    ws = wb.create_sheet("Settings")
    s = ex.settings
    rows = [("Key", "Value"), ("format", MARKER), ("format_version", WORKBOOK_VERSION),
            ("source", ex.source), ("value_label", ex.value_label),
            ("value_mode", s.value_mode), ("pixel_size_um", s.pixel_size_um),
            ("settings_json", s.to_json())]
    if ex.color_scale is not None:
        rows.append(("color_scale_units", ex.color_scale.units))
    for r in rows:
        _append(ws, r)
    header(ws)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 120
    for ws in wb:
        ws.sheet_view.showGridLines = ws.title not in ("Summary",)
    wb.save(path)
    return str(Path(path))


# ------------------------------------------------------------------ reading back
@dataclass
class ExtractionTable:
    """An extraction workbook read back: the profile table and its settings."""
    profile: pd.DataFrame
    settings: Optional[ExtractionSettings]
    source: str
    value_label: str
    pixel_size_um: Optional[float]
    path: str = ""

    @property
    def statistics(self) -> List[str]:
        return [c for c in ("Clean_Mean", "Clean_Median", "Raw_Mean") if c in self.profile]


def is_extraction_workbook(path) -> bool:
    p = Path(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        try:
            if "Settings" not in wb.sheetnames:
                return False
            for row in wb["Settings"].iter_rows(min_row=1, max_row=4, values_only=True):
                if row and row[0] == "format" and row[1] == MARKER:
                    return True
            return False
        finally:
            wb.close()
    except Exception:
        return False


def read_extraction(path) -> ExtractionTable:
    p = Path(path)
    prof = pd.read_excel(p, sheet_name="Profile")
    kv = pd.read_excel(p, sheet_name="Settings")
    meta = dict(zip(kv.iloc[:, 0].astype(str), kv.iloc[:, 1]))
    if meta.get("format") != MARKER:
        raise ValueError(f"{p.name} is not a Diffusor image-profile workbook")
    settings = None
    if isinstance(meta.get("settings_json"), str):
        try:
            settings = ExtractionSettings.from_json(meta["settings_json"])
        except Exception:
            settings = None
    px = meta.get("pixel_size_um")
    px = float(px) if px is not None and not (isinstance(px, float) and math.isnan(px)) else None
    return ExtractionTable(prof, settings, str(meta.get("source", "")),
                           str(meta.get("value_label", "grey value")), px, str(p))


def table_from_extraction(ex: ProfileExtraction) -> ExtractionTable:
    """The same object :func:`read_extraction` returns, without a round trip to disk."""
    return ExtractionTable(ex.profile_table(), ex.settings, ex.source, ex.value_label,
                           ex.settings.pixel_size_um)


def composition_table(table: ExtractionTable, pixel_size_um: Optional[float],
                      statistic: str = "Clean_Mean", uncertainty: Optional[str] = "Clean_SE",
                      calibration: Optional[GreyscaleCalibration] = None,
                      name: str = "Composition") -> pd.DataFrame:
    """Distance in micrometres and composition with its uncertainty, ready to model.

    ``calibration`` maps the profile values (grey, channel or legend values) to
    composition; see :func:`greyscale.calibrate`. ``None`` means the values
    already are compositions, as in a quantitative map or a legend drawn in wt%.
    The uncertainty combines the scatter across lines with the calibration's
    own uncertainty, through :func:`greyscale.apply_calibration`.
    """
    prof = table.profile
    if pixel_size_um is None or not pixel_size_um > 0:
        raise ValueError("a pixel size (um per pixel) is needed to turn pixels into distance")
    if statistic not in prof:
        raise ValueError(f"the workbook has no column {statistic}")
    x = np.asarray(prof["Distance_px"], float) * float(pixel_size_um)
    off = np.asarray(prof["Offset_px"], float) * float(pixel_size_um)
    g = np.asarray(prof[statistic], float)
    scatter = None if not uncertainty else np.asarray(prof[uncertainty], float)
    if calibration is None:
        C, sig = g, scatter
    else:
        _, C, sig = apply_calibration(x, g, calibration, grey_scatter=scatter)
    n_col = "Clean_N" if statistic.startswith("Clean") else "Raw_N"
    out = pd.DataFrame({"Distance_um": x, "Offset_um": off, name: C})
    if sig is not None:
        out[name + "_err"] = sig
    out["Image_" + statistic] = g
    if scatter is not None:
        out["Image_" + uncertainty] = scatter
    if n_col in prof:
        out["Lines_used"] = np.asarray(prof[n_col])
    return out[np.isfinite(out[name])].reset_index(drop=True)
