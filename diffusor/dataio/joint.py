"""Export a shared-duration study and each constituent profile's provenance."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from .export import result_dict, save_results
from .workbook import _table
from .figures import joint_figure


def save_joint_results(directory, result):
    out = Path(directory)
    out.mkdir(parents=True, exist_ok=True)
    payload = {"model_family": "independent_scalar_shared_duration", "t_seconds": result.t_seconds,
               "chi2": result.chi2, "dof": result.dof, "reduced_chi2": result.reduced_chi2,
               "warnings": result.warnings,
               "profiles": [{"name": name, **result_dict(fit)} for name, fit in zip(result.names, result.profiles)]}
    json_path = out / 'joint_results.json'
    json_path.write_text(json.dumps(payload, indent=2), encoding='utf-8')
    wb = Workbook()
    ws = wb.active
    ws.title = "Joint results"
    _table(ws, ("Property", "Value"), [("Shared duration (s)", result.t_seconds),
           ("Total chi-squared", result.chi2), ("Joint degrees of freedom", result.dof),
           ("Reduced chi-squared", result.reduced_chi2),
           *[("Assumption", w) for w in result.warnings]])
    from openpyxl.styles import Alignment
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 115
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(row[1].value, str):
            ws.row_dimensions[row[1].row].height = 48
    summary = wb.create_sheet("Profiles")
    _table(summary, ("Profile", "Coefficient", "Points", "Chi-squared contribution", "RMSE (profile units)", "Detailed report folder"),
           [(name, fit.model.coefficient.key, len(fit.x_data), fit.stats.chi2, fit.stats.rmse, f'profile_{i:03}')
            for i, (name, fit) in enumerate(zip(result.names, result.profiles), 1)])
    for col in 'ABCDEF':
        summary.column_dimensions[col].width = 32
    summary.column_dimensions['B'].width = 48
    for i, fit in enumerate(result.profiles, 1):
        save_results(out / f'profile_{i:03}', fit)
        data = wb.create_sheet(f"Profile {i:03}")
        _table(data, ("Distance (um)", "Measured", "Model", "Sigma", "Standardised residual"),
               zip(fit.x_data, fit.C_data, fit.C_model, fit.sigma, (fit.C_data-fit.C_model)/fit.sigma))
        for col in 'ABCDE':
            data.column_dimensions[col].width = 25
    wb.save(out / 'joint_results.xlsx')
    pd.DataFrame({"time_s": result.scan_times, "chi2": result.scan_chi2}).to_csv(out / 'joint_time_scan.csv', index=False)
    fig = joint_figure(result)
    fig.savefig(out / 'joint_profiles.png', dpi=220, bbox_inches='tight')
    fig.savefig(out / 'joint_profiles.svg', bbox_inches='tight')
    return {"json": str(json_path), "workbook": str(out / 'joint_results.xlsx'), "figure": str(out / 'joint_profiles.png')}
