"""Figures and exports for multicomponent (garnet) and isotope studies."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.styles import Alignment

from .. import __version__
from ..references import get as reference
from ..thermo.units import human_time
from .figures import GOLD, INK, MUTED, TEAL, _style
from .workbook import _table

COLOURS = {"Fe": "#7a3b2e", "Mg": TEAL, "Mn": "#6a4c93", "Ca": GOLD}


def _panel(ax, res_ax, x, y, sigma, model, initial=None, title="", ylabel="", colour=TEAL):
    _style(ax)
    _style(res_ax)
    order = np.argsort(x)
    x, y, model = x[order], y[order], model[order]
    sigma = None if sigma is None else np.broadcast_to(sigma, y.shape)[order]
    if initial is not None:
        xi, ci = initial
        ax.plot(xi, ci, color=MUTED, linestyle="--", linewidth=1.0, label="Initial state")
    ax.plot(x, model, color=colour, linewidth=2.4, label="Model")
    ax.errorbar(x, y, yerr=sigma, fmt="o", markersize=3.5, color=INK, markerfacecolor="white",
                markeredgewidth=.9, ecolor="#aab8bf", elinewidth=.7, capsize=1.2, label="Measured")
    ax.set_title(title, loc="left", fontsize=11, color=INK, fontweight="bold")
    ax.set_ylabel(ylabel, color=INK)
    ax.tick_params(labelbottom=False)
    r = (y - model) / sigma if sigma is not None else y - model
    if sigma is not None:
        res_ax.axhspan(-2, 2, color="#edf5f5", zorder=0)
    res_ax.axhline(0, color=MUTED, linewidth=.8)
    res_ax.plot(x, r, "o", color=GOLD, markersize=3)
    res_ax.set_ylabel("Residual / σ" if sigma is not None else "Residual", color=INK)
    res_ax.set_xlabel("Distance (µm)", color=INK)


def multicomponent_figure(result) -> Figure:
    comps = list(result.fitted)
    n = len(comps)
    fig = Figure(figsize=(10, 2.7 * n + 1.0), facecolor="#fafcfb", layout="constrained")
    FigureCanvasAgg(fig)
    fig.suptitle(f"DIFFUSOR  /  MULTICOMPONENT  ({result.model.family.mineral})\n"
                 f"{human_time(result.t_seconds)} · χ² / dof = {result.stats.reduced_chi2:.3g} · "
                 f"{reference(result.model.family.citation).short()}", x=.06, ha="left", color=INK, fontsize=13)
    grid = fig.add_gridspec(2 * n, 1, height_ratios=[3, 1] * n)
    xi = np.linspace(result.x.min(), result.x.max(), 501)
    init = result.model.initial_state(xi)
    for k, c in enumerate(comps):
        ax = fig.add_subplot(grid[2 * k])
        res = fig.add_subplot(grid[2 * k + 1], sharex=ax)
        _panel(ax, res, result.x, result.data[c], result.sigma.get(c), result.prediction[c],
               (xi, init[c]), f"{c}   ({result.model.family.composition_labels.get(c, c)})", f"X_{c}",
               COLOURS.get(c, TEAL))
        if k == 0:
            ax.legend(loc="best", frameon=False, fontsize=8)
    return fig


def isotope_figure(result) -> Figure:
    keys = list(result.observed)
    n = len(keys)
    fig = Figure(figsize=(10, 2.8 * n + 1.0), facecolor="#fafcfb", layout="constrained")
    FigureCanvasAgg(fig)
    extra = f" · β fitted {result.beta:.3g}" if result.beta is not None else ""
    fig.suptitle(f"DIFFUSOR  /  ISOTOPE STUDY\n{human_time(result.t_seconds)} · "
                 f"χ² / dof = {result.stats.reduced_chi2:.3g}{extra}", x=.06, ha="left", color=INK, fontsize=13)
    grid = fig.add_gridspec(2 * n, 1, height_ratios=[3, 1] * n)
    for k, key in enumerate(keys):
        ax = fig.add_subplot(grid[2 * k])
        res = fig.add_subplot(grid[2 * k + 1], sharex=ax)
        label = result.labels.get(key, key)
        _panel(ax, res, result.x, result.observed[key], result.sigma[key], result.predicted[key],
               None, label, label + (" (‰)" if key not in ("C", "XFe") else ""),
               GOLD if key not in ("C", "XFe") else TEAL)
        if k == 0:
            ax.legend(loc="best", frameon=False, fontsize=8)
    return fig


def _write_book(path, summary_rows, sheets):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    _table(ws, ("Property", "Value"), summary_rows)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for name, header, rows in sheets:
        s = wb.create_sheet(name[:31])
        _table(s, header, rows)
        for i in range(len(header)):
            s.column_dimensions[chr(65 + i)].width = 22
    wb.save(path)


def save_multicomponent_results(directory, result) -> dict:
    out = Path(directory)
    out.mkdir(parents=True, exist_ok=True)
    m, fam = result.model, result.model.family
    methods = (fam.methods_sentence() + f" The duration was fitted jointly to the {', '.join(result.fitted)} "
               f"profiles by weighted least squares with Diffusor {__version__}. " + m.describe() + ".")
    payload = {"model_family": "multicomponent_ideal_ionic", "diffusor_version": __version__,
               "family": fam.key, "citation": fam.citation, "t_seconds": result.t_seconds,
               "free": result.free, "fitted": list(result.fitted), "dependent": m.dependent,
               "T_K": m.T_K, "P_Pa": m.P_Pa, "options": fam._options(m.options),
               "geometry": m.geometry.kind, "left_fixed": m.left_fixed, "right_fixed": m.right_fixed,
               "chi2": result.stats.chi2, "reduced_chi2": result.stats.reduced_chi2, "dof": result.stats.dof,
               "per_component_chi2": {c: s.chi2 for c, s in result.per_component.items()},
               "monte_carlo_times_s": None if result.mc_times is None else result.mc_times.tolist(),
               "family_description": fam.describe(), "warnings": result.warnings, "methods": methods}
    (out / "multicomponent_results.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
    rows = [("Duration (s)", result.t_seconds), ("Duration", human_time(result.t_seconds)),
            ("Tracer family", f"{fam.label} [{fam.key}]"), ("Source", reference(fam.citation).full()),
            ("Model", m.describe()), ("Chi-squared", result.stats.chi2),
            ("Reduced chi-squared", result.stats.reduced_chi2), ("Methods", methods),
            *[("Caveat", w) for w in result.warnings]]
    if result.mc_times is not None and result.mc_times.size:
        lo, med, hi = np.percentile(result.mc_times, [16, 50, 84])
        rows[2:2] = [("Monte Carlo median (s)", med), ("Monte Carlo 16th percentile (s)", lo),
                     ("Monte Carlo 84th percentile (s)", hi), ("Monte Carlo draws", int(result.mc_times.size))]
    comps = fam.components
    data_rows = [(xv, *[result.data[c][i] if c in result.data else None for c in comps],
                  *[result.prediction[c][i] for c in comps]) for i, xv in enumerate(result.x)]
    header = ("Distance (um)", *[f"{c} measured" for c in comps], *[f"{c} model" for c in comps])
    _write_book(out / "multicomponent_results.xlsx", rows, [("Profiles", header, data_rows)])
    pd.DataFrame({"time_s": result.scan_times, "chi2": result.scan_chi2}).to_csv(out / "time_scan.csv", index=False)
    pd.DataFrame(data_rows, columns=header).to_csv(out / "profiles.csv", index=False)
    (out / "methods.txt").write_text(methods + "\n", encoding="utf-8")
    fig = multicomponent_figure(result)
    fig.savefig(out / "multicomponent_profiles.png", dpi=220, bbox_inches="tight")
    fig.savefig(out / "multicomponent_profiles.svg", bbox_inches="tight")
    return {"json": str(out / "multicomponent_results.json"), "workbook": str(out / "multicomponent_results.xlsx")}


def save_isotope_results(directory, result, description: str = "") -> dict:
    out = Path(directory)
    out.mkdir(parents=True, exist_ok=True)
    payload = {"model_family": "isotope_fractionation", "diffusor_version": __version__,
               "t_seconds": result.t_seconds, "beta_fitted": result.beta,
               "profiles": list(result.observed), "chi2": result.stats.chi2,
               "reduced_chi2": result.stats.reduced_chi2, "dof": result.stats.dof,
               "per_profile_chi2": {k: s.chi2 for k, s in result.per_profile.items()},
               "model": description, "warnings": result.warnings}
    (out / "isotope_results.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
    keys = list(result.observed)
    header = ("Distance (um)", *[f"{result.labels.get(k, k)} measured" for k in keys],
              *[f"{result.labels.get(k, k)} sigma" for k in keys], *[f"{result.labels.get(k, k)} model" for k in keys])
    data_rows = [(xv, *[result.observed[k][i] for k in keys], *[result.sigma[k][i] for k in keys],
                  *[result.predicted[k][i] for k in keys]) for i, xv in enumerate(result.x)]
    rows = [("Duration (s)", result.t_seconds), ("Duration", human_time(result.t_seconds)),
            ("Model", description), ("Chi-squared", result.stats.chi2),
            ("Reduced chi-squared", result.stats.reduced_chi2),
            *([("Fitted beta", result.beta)] if result.beta is not None else []),
            *[("Caveat", w) for w in result.warnings]]
    _write_book(out / "isotope_results.xlsx", rows, [("Profiles", header, data_rows)])
    pd.DataFrame(data_rows, columns=header).to_csv(out / "profiles.csv", index=False)
    pd.DataFrame({"time_s": result.scan_times, "chi2": result.scan_chi2}).to_csv(out / "time_scan.csv", index=False)
    fig = isotope_figure(result)
    fig.savefig(out / "isotope_profiles.png", dpi=220, bbox_inches="tight")
    fig.savefig(out / "isotope_profiles.svg", bbox_inches="tight")
    return {"json": str(out / "isotope_results.json"), "workbook": str(out / "isotope_results.xlsx")}
