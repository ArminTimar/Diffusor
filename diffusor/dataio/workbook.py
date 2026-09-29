"""Native Excel result reports using Diffusor's existing openpyxl dependency.

Workbooks are numerical snapshots of a completed fit, not a second diffusion
solver. Measured values, fitted values and uncertainties retain numeric types.
"""
from pathlib import Path
import math

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INK = "243746"
TEAL = "007F82"
LIGHT = "EDF5F5"
GOLD = "D39034"


def _value(v):
    if isinstance(v, (np.integer, np.floating)):
        v = v.item()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _append(ws, values):
    ws.append([_value(v) for v in values])
    # Sources and filenames are text, even if they start with '='.
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _table(ws, headings, rows):
    _append(ws, headings)
    for row in rows:
        _append(ws, row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=INK)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10, color=INK)
            c.alignment = Alignment(vertical="center")
            if isinstance(c.value, (int, float)):
                c.number_format = "0.000000E+00"
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)


def _chart(ws, ycols, title, ytitle):
    chart = ScatterChart()
    chart.title = title
    chart.x_axis.title = "Distance (µm)"
    chart.y_axis.title = ytitle
    chart.width, chart.height = 23, 10
    chart.style = 13
    xref = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    for col, color, markers in ycols:
        series = Series(Reference(ws, min_col=col, min_row=1, max_row=ws.max_row), xref,
                        title_from_data=True)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 23000
        if markers:
            series.marker.symbol = "circle"
            series.marker.size = 4
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.noFill = True
        chart.series.append(series)
    return chart


def save_workbook(path, fit_result, profile_frame, metadata, methods, mc_result=None):
    """Write a readable, filterable, charted .xlsx snapshot and return its path."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Results"
    summary.sheet_properties.tabColor = TEAL
    for row in ((), ("Diffusor results",), ("Numerical snapshot of the completed fit",), (),
                ("Best-fit time (s)", fit_result.t_seconds),
                ("Best-fit time (years)", fit_result.t_years),
                ("Nominal temperature (°C)" if fit_result.model.history is not None else "Temperature (°C)", metadata["conditions"]["T_C"]),
                ("Pressure (MPa)", metadata["conditions"]["P_Pa"] / 1e6),
                ("Coefficient", fit_result.model.coefficient.key),
                ("Transport kind", fit_result.model.coefficient.kind),
                ("Validation", fit_result.model.coefficient.validation_level),
                ("RMSE (profile units)", fit_result.stats.rmse),
                ("Reduced chi-squared", fit_result.stats.reduced_chi2),
                ("Solver", fit_result.route)):
        _append(summary, row)
    if mc_result is not None:
        for name, value in (("MC median (s)", mc_result.median),
                            ("MC 16th percentile (s)", mc_result.p16),
                            ("MC 84th percentile (s)", mc_result.p84),
                            ("MC failed draws", mc_result.n_failed), ("MC seed", mc_result.seed)):
            _append(summary, (name, value))
    _append(summary, ())
    _append(summary, ("Interpretation and limitations",))
    warnings = list(dict.fromkeys(metadata["warnings"] +
                                 (list(mc_result.warnings) if mc_result is not None else [])))
    for warning in warnings:
        _append(summary, ("Caveat", warning))
        summary.cell(summary.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[summary.max_row].height = max(45, 15 * math.ceil(len(warning)/70))
    summary.column_dimensions["A"].width = 29
    summary.column_dimensions["B"].width = 72
    summary.column_dimensions["C"].width = 3
    for row in summary:
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000E+00"
    summary["A2"].font = Font(name="Arial", size=16, bold=True, color=TEAL)
    summary["A5"].font = Font(name="Arial", size=11, bold=True, color=INK)
    summary["B5"].font = Font(name="Arial", size=14, bold=True, color=TEAL)

    profiles = wb.create_sheet("Profile")
    ordered = profile_frame.sort_values("x_um")
    _table(profiles, list(ordered.columns), ordered.itertuples(index=False, name=None))
    for i in range(1, profiles.max_column+1):
        profiles.column_dimensions[get_column_letter(i)].width = 23
    summary.add_chart(_chart(profiles, [(2, INK, True), (3, TEAL, False)],
                             "Measured profile and fitted model", "Concentration (input units)"), "D4")
    summary.add_chart(_chart(profiles, [(4, GOLD, True)],
                             "Residuals", "Measured − model"), "D25")
    coef = wb.create_sheet("Coefficient")
    c = fit_result.model.coefficient
    _table(coef, ("Property", "Value"),
           [("Key", c.key), ("Source", metadata["citations"][c.citation]),
            ("Equation", c.equation_text), ("Source location", c.equation_number),
            ("Verified from", c.verified_from), ("Reference state", c.reference_state),
            ("Transport kind", c.kind), ("State variable", c.transported_variable),
            ("Uncertainty", c.uncertainty_note),
            *[(k, str(getattr(c, k))) for k in ("T_range", "P_range", "fo2_range", "X_range")],
            *[(p.name, f"{p.value:g} {p.unit}; " + (f"reported uncertainty {p.sigma:g} ({p.sigma_level})" if p.sigma > 0 else "parameter uncertainty not sampled"))
              for p in c.params.values()],
            *[("Limitation", s) for s in c.calibration_notes]])
    coef.column_dimensions["A"].width = 28
    coef.column_dimensions["B"].width = 115
    for row in coef.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        coef.row_dimensions[row[1].row].height = max(25, 15 * math.ceil(len(str(row[1].value))/110))
    meth = wb.create_sheet("Methods")
    _table(meth, ("Methods and references",), [(line,) for line in methods.splitlines()])
    meth.column_dimensions["A"].width = 145
    for row in meth.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        meth.row_dimensions[row[0].row].height = max(16, 15 * math.ceil(len(str(row[0].value or ''))/140))
    if mc_result is not None:
        draws = wb.create_sheet("Monte Carlo")
        _table(draws, ("Successful draw", "Time (s)", "Time (years)"),
               [(i+1, t, t / 31557600) for i, t in enumerate(mc_result.times)])
        for col in "ABC":
            draws.column_dimensions[col].width = 24
    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
    wb.save(path)
    return str(Path(path))
