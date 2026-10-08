"""Recompute published timescale tables from their own inputs and laws.

Run from the repository root: python examples/validation/check_timescale_tables.py [source_dir]
Reads the downloaded supplementary workbooks (openpyxl), like extract.py.

* Chamberlain et al. (2014), Electronic Appendix 7: the '+ timescale (-30 C)' and
  '- timescale (+30 C)' columns of the Sr, Ba and Ti sheets are the best-fit time
  scaled by D(T)/D(T -/+ 30 K), with the Table 1 laws (p. 5): Sr D0 8.4 m2/s,
  450 kJ/mol; Ba 0.29 m2/s, 455 kJ/mol; Ti 7e-8 m2/s, 273 kJ/mol.
* Petrone et al. (2018), Table S4: t = (sqrt(4Dt))^2 / (4 D(T)) with
  D = 9.5e-5 m2/s exp(-406 kJ/mol / RT), the Dimanov & Sautter (2000) values as
  printed in Supplementary Material 1 (p. 13), for rows fitted with the
  semi-infinite erf (no finite-reservoir width h).

These check the arithmetic between columns, not a refit of profiles: neither
supplement tabulates the measured traverses.
"""
from pathlib import Path
import json
import sys

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('papers/Supplementaries and data')
R = 8.314462618
YEAR = 365.25 * 86400
number = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)


def D(D0, E, T_K):
    return D0 * np.exp(-E / (R * T_K))


def chamberlain():
    p = SOURCE/'chamberlain2014/410_2014_1034_MOESM7_ESM.xlsx'
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    laws = {'Sr in sanidine': (8.4, 450e3, 4), 'Ba in sanidine': (0.29, 455e3, 3), 'Ti in quartz': (7e-8, 273e3, 3)}
    out = {}
    for sheet, (D0, E, it) in laws.items():
        errs, n = [], 0
        for i, r in enumerate(wb[sheet].iter_rows(values_only=True), 1):
            if i == 1 or not all(number(r[k]) for k in (1, it, it+1, it+2)):
                continue
            T, t = r[1], r[it]
            plus, minus = t*D(D0, E, T)/D(D0, E, T-30), t*D(D0, E, T)/D(D0, E, T+30)
            errs += [abs(plus/r[it+1]-1), abs(minus/r[it+2]-1)]
            n += 1
        out[sheet] = dict(rows=n, max_relative_difference=max(errs))
    return out


def petrone2018():
    p = SOURCE/'petrone2018/1-s2.0-S0012821X18301936-mmc2.xlsx'
    rr = list(openpyxl.load_workbook(p, read_only=True, data_only=True)['S4'].values)
    # (T, sqrt(4Dt), h, time) columns of the three boundary blocks of Table S4
    blocks = [(8, 9, 12, 14), (23, 24, 27, 29), (38, 39, 42, 44)]
    checked, finite, rows = 0, 0, []
    for i, r in enumerate(rr, 1):
        if i < 4:
            continue
        for b, (iT, iw, ih, itime) in enumerate(blocks, 1):
            if not (number(r[iT]) and number(r[iw]) and number(r[itime])):
                continue
            if number(r[ih]):
                finite += 1
                continue
            w = r[iw]*1e-6
            t = w**2/(4*D(9.5e-5, 406e3, r[iT]+273.15))/YEAR
            checked += 1
            rows.append(dict(row=i, crystal=str(r[2]), block=b, T_C=r[iT], sqrt4Dt_um=r[iw],
                             printed_years=r[itime], recomputed_years=t, ratio=t/r[itime]))
    ratios = np.array([v['ratio'] for v in rows])
    within = [v for v in rows if abs(v['ratio']-1) <= .05 or abs(v['recomputed_years']-v['printed_years']) <= .01]
    return dict(rows_checked=checked, finite_reservoir_rows_skipped=finite,
                within_5_percent_or_rounding=len(within), median_ratio=float(np.median(ratios)),
                outside=[v for v in rows if v not in within])


def main():
    result = dict(chamberlain2014=chamberlain(), petrone2018=petrone2018())
    (ROOT/'timescale_tables_results.json').write_text(json.dumps(result, indent=2, ensure_ascii=False)+'\n',
                                                     encoding='utf8')
    print(json.dumps({k: {kk: vv for kk, vv in v.items() if kk != 'outside'} for k, v in result.items()}, indent=1))
    for v in result['petrone2018']['outside']:
        print(v)


if __name__ == '__main__':
    main()
