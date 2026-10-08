"""Extract downloaded primary data without digitising or inventing measurements.

Run from the repository root: python examples/validation/extract.py [source_dir]
Requires openpyxl; Ruth's legacy workbook additionally requires xlrd.
Original workbooks are read only. Row numbers and SHA256 hashes are retained.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import re
import sys

import openpyxl

OUT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('papers/Supplementaries and data')
RECORDS = []


def number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def rows(path, sheet):
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(w[sheet].values)
    finally:
        w.close()


def sheets(path):
    w = openpyxl.load_workbook(path, read_only=True)
    result = w.sheetnames
    w.close()
    return result


def emit(study, sample, path, sheet, data, mapping, scale=1, setup=None, notes='', fo_scale=1):
    """data: (Excel row, row values); mapping: output name -> zero-based column."""
    selected = []
    for rownum, r in data:
        if not all(number(r[mapping[k]]) for k in ('Distance_um', 'Fo_mol')):
            continue
        d = {k: r[i] if number(r[i]) else '' for k, i in mapping.items()}
        d['Distance_um'] *= scale
        for k in ('Fo_mol','Initial_Fo_mol','Fo_sigma'):
            if k in d and number(d[k]):
                d[k] *= fo_scale
        d['source_row'] = rownum
        selected.append(d)
    if len(selected) < 3:
        return
    key = study + '_' + re.sub(r'[^a-z0-9]+', '_', str(sample).lower()).strip('_')
    dest = OUT / 'profiles' / (key + '.csv')
    dest.parent.mkdir(exist_ok=True)
    with dest.open('w', newline='', encoding='utf8') as f:
        writer = csv.DictWriter(f, fieldnames=list(mapping) + ['source_row'])
        writer.writeheader()
        writer.writerows(selected)
    RECORDS.append(dict(key=key, study=study, sample=str(sample), file='profiles/'+dest.name,
                        source_file=path.relative_to(SOURCE).as_posix(), source_sheet=sheet,
                        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                        n_points=len(selected), setup=setup, notes=notes))


def grouped(study, path, sheet, id_column, mapping, start=1, scale=1, transform=str, notes='', fo_scale=1):
    groups = {}
    for i, r in enumerate(rows(path, sheet), 1):
        identity = ' '.join(str(r[c]) for c in id_column) if isinstance(id_column,tuple) else r[id_column]
        if i < start or not isinstance(identity, str):
            continue
        groups.setdefault(transform(identity), []).append((i, r))
    for name, data in groups.items():
        emit(study, name, path, sheet, data, mapping, scale=scale, notes=notes,fo_scale=fo_scale)


def extract():
    # All 19 archived Lynn traverses, including the two-event Ol 8.
    p = SOURCE/'lynn_2024_data/3087-1_Lynn_2020Kilauea_DataTables_01122024.xlsx'
    for sheet in sheets(p):
        if not sheet.startswith('Ol '):
            continue
        rr = rows(p, sheet)
        setup = dict(T_C=rr[5][15], P_MPa=rr[5][16], log_fo2_bar=rr[5][17],
                     angles_deg=list(rr[9][15:18]), published_days=rr[13][15],
                     published_sigma_days=rr[13][16], coefficient='ol_FeMg_dohmen_chakraborty2007_tamed')
        if not all(number(setup[k]) for k in ('T_C','P_MPa','log_fo2_bar','published_days')):
            setup = None
        emit('lynn2024', sheet, p, sheet, enumerate(rr[3:],4),
             dict(Distance_um=1,Fo_mol=10,Initial_Fo_mol=12,Published_model_Fo_mol=13,
                  FeO_wt=3,MgO_wt=4,CaO_wt=5,NiO_wt=7), setup=setup,
             notes='Workbook inputs retained, including 42 MPa versus 45 MPa in the article. '
                   'Initial column is sampled on the EPMA positions. The sub-grid interface and original solver are not archived.')
    p = SOURCE/'mutch2019/Data_S1.xlsx'
    grouped('mutch2019',p,sheets(p)[0],0,
            dict(Distance_um=1,Fo_mol=2,Fo_sigma=3,Ni_ppm=8,Ni_sigma=9,Mn_ppm=10,
                 Mn_sigma=11,Initial_Fo_mol=25,Initial_Ni_ppm=26,Initial_Mn_ppm=27),start=2,fo_scale=100,
            notes='Source Fo, Fo_sigma and initial Fo mole fractions multiplied by 100 to mol%. '
                  'Joint Bayesian Fe-Mg/Ni/Mn inversion with study-specific D regressions; not a TaMED replication.')
    p = next((SOURCE/'mourey2023').glob('*.xlsx'))
    for sheet in sheets(p):
        if not sheet.startswith('KE62-'):
            continue
        rr = rows(p,sheet)
        # Individual sample sheets have different analytical routines/columns.
        header = [str(v).strip().lower() for v in rr[2]]
        find = lambda s: next(i for i,v in enumerate(header) if s in v)
        grouped('mourey2023',p,sheet,0,dict(Distance_um=find('reldist'),Fo_mol=find('fo'),
                FeO_wt=find('feo'),MgO_wt=find('mgo'),CaO_wt=find('cao'),NiO_wt=find('nio')),start=4,
                notes='Profile includes possible multiple events. Rim positions are constrained from Ca and images in the study; no numerical initial arrays in this workbook.')
    p = SOURCE/'sundermeyer2019/410_2019_1642_MOESM1_ESM.xlsx'
    for sheet in sheets(p):
        grouped('sundermeyer2020_reunion',p,sheet,0,
                dict(Distance_um=1,Fo_mol=17,FeO_wt=8,MgO_wt=10,CaO_wt=6,NiO_wt=14),start=7,
                notes='Whole measured traverse retained. The paper excludes the late outermost rim (<3 um) and fits only selected gradients; exact per-profile fit masks are not tabulated.')
    p = SOURCE/'sundermeyer2020/410_2020_1715_MOESM2_ESM.xlsx'
    for sheet in sheets(p):
        grouped('sundermeyer2020_eifel',p,sheet,1,
                dict(Distance_um=0,Fo_mol=25,FeO_wt=12,MgO_wt=10,CaO_wt=14,NiO_wt=18),start=7,
                transform=lambda x:re.sub(r'^Line\d+','',x),
                notes='Crystal names retain traverse suffixes. Growth zones and cooling/effective temperatures need to be matched before a published-time replication.')
    p = SOURCE/'gordeychik2018/41598_2018_30133_MOESM4_ESM.xlsx'
    grouped('gordeychik2018',p,'Table SM2-A',(3,4,5),
            dict(Distance_um=12,Fo_mol=25,FeO_wt=15,MgO_wt=14,CaO_wt=20,NiO_wt=18),start=3,scale=1000,
            notes='Source distances are mm, converted to um. Growth plus diffusion interpretation; retain whole traverse without assuming a sharp-step initial state.')
    p = SOURCE/'ruprecht_plank2023/41586_2013_BFnature12342_MOESM46_ESM.xlsx'
    for sheet in sheets(p):
        if not sheet.startswith('IZ-'):
            continue
        emit('ruprecht2013',sheet,p,sheet,enumerate(rows(p,sheet)[5:],6),
             dict(Distance_um=21,Fo_mol=39,Ni_ppm=36),
             notes='San Carlos-normalised LA-ICPMS values (columns V, AK, AN). '
                   'Study is 2013, despite the local folder name. Full composition/fO2-dependent Ni model is not the registry Fo90 fit.')
    import xlrd
    p = next((SOURCE/'ruth2018').glob('*.xls'))
    w = xlrd.open_workbook(str(p)); s = w.sheet_by_name('3 Data')
    groups = {}
    for i in range(6,s.nrows):
        r=s.row_values(i)
        if r[0]:groups.setdefault(r[0],[]).append((i+1,r))
    for name,data in groups.items():
        emit('ruth2018',name,p,s.name,data,dict(Distance_um=7,Fo_mol=18,FeO_wt=11,MgO_wt=13,CaO_wt=14,NiO_wt=16),
             notes='Measured rim-to-interior profiles; archived composition table does not include per-traverse initial conditions or EBSD angles.')
    extract_lynn_mauna_loa()
    extract_iovine()
    extract_sato()
    extract_morgado()
    (OUT/'manifest.json').write_text(json.dumps(RECORDS,indent=2,ensure_ascii=False)+'\n',encoding='utf8')
    from collections import Counter
    print(dict(Counter(r['study'] for r in RECORDS)))


def emit_table(study, sample, path, sheet, points, *, mineral, species, value_column, spec,
               setup=None, notes='', extra=None):
    """Write one traverse of a non-olivine (or non-Lynn-layout) study.

    points: list of dicts holding the CSV columns, each with 'source_row' (an Excel
    row number) or 'source_cell' (an A1 reference, for transposed sheets).
    """
    if len(points) < 3:
        return
    key = study + '_' + re.sub(r'[^a-z0-9]+', '_', str(sample).lower()).strip('_')
    dest = OUT / 'profiles' / (key + '.csv')
    dest.parent.mkdir(exist_ok=True)
    fields = list(dict.fromkeys(k for p in points for k in p))
    with dest.open('w', newline='', encoding='utf8') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(points)
    RECORDS.append(dict(key=key, study=study, sample=str(sample), file='profiles/'+dest.name,
                        source_file=path.relative_to(SOURCE).as_posix(), source_sheet=sheet,
                        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                        n_points=len(points), mineral=mineral, species=species,
                        value_column=value_column, spec=spec, setup=setup, notes=notes,
                        **(extra or {})))


def extract_lynn_mauna_loa():
    """Lynn et al. Mauna Loa 2022 workbook: olivine traverses with the workbook's own conditions.

    The workbook names no diffusion law and the article is not in the local source set,
    so no model setup is attached: T, P, fO2, angles and the S6 timescale are kept as
    source values only.
    """
    p = SOURCE/'lynn_2024_data/3447-1_Lynn_ML2022_SupplementaryData1_Revision1.xlsx'
    if not p.exists():
        return
    times = {}
    for i, r in enumerate(rows(p, 'S6. Timescales'), 1):
        if i > 2 and isinstance(r[2], str) and r[2].startswith('ORI') and number(r[6]):
            times[f'{r[2]}_ol{int(r[3])}'] = dict(sample=r[0], location=r[1], core_Fo=r[4], rim_Fo=r[5],
                median_days=r[6], day_erupted=r[7], adjusted_median_days=r[8],
                adjusted_min95_days=r[9], adjusted_max95_days=r[10], source_row=i)
    for sheet in sheets(p):
        if not re.fullmatch(r'ORI-\d_ol\d+', sheet):
            continue
        rr = rows(p, sheet)
        head = [str(v).strip() if v is not None else '' for v in rr[1]]
        col = lambda name: head.index(name)
        iD, iF, iV = col('Dist'), col('Fo'), col('Valid')
        iT = col('T (°C)')
        first = rr[2]
        conditions = dict(T_C=first[iT], P_MPa=first[iT+1], log_fo2_bar=first[iT+2],
                          angles_deg=list(first[iT+3:iT+6]))
        points = []
        for n, r in enumerate(rr[2:], 3):
            if number(r[iD]) and number(r[iF]):
                points.append(dict(Distance_um=r[iD], Fo_mol=r[iF],
                                   Valid=r[iV] if number(r[iV]) else '',
                                   SiO2_wt=r[1], MgO_wt=r[2], CaO_wt=r[3], MnO_wt=r[4],
                                   FeO_wt=r[5], NiO_wt=r[6], source_row=n))
        emit_table('lynn_ml2022', sheet, p, sheet, points, mineral='olivine', species='Fe-Mg',
                   value_column='Fo_mol',
                   spec=dict(distance_column='Distance_um', column_a='Fo_mol', mode='A',
                             distance_unit='um'),
                   notes='Workbook T, P, log fO2 and EBSD angles and the S6 timescale are kept as '
                         'source values. The workbook names no diffusion law and the article is not '
                         'in the local source set, so no model preset is attached. Valid = 0 marks '
                         'points the workbook flags. The Dist column carries no unit (micrometres '
                         'by its 10-unit spacing).',
                   extra=dict(source_conditions=conditions, source_timescale=times.get(sheet)))


# Iovine et al. (2017) Tables 1 (BaO and greyscale) and 2 (X-ray), 930 C, read from the
# rendered PDF pages 10-11. Key: (sheet, profile kind, Line, Profile) of the supplement
# workbook. 'L1 2°' and 'L1 2nd' are read as the second profile of line 1 (Diffusor's
# reading; the paper does not define the suffix).
IOVINE_PUBLISHED = {
    ('A CX4', 'grey', 2, None): ('Table 1', 'cx4_L2', 16),
    ('A CX5', 'grey', 1, None): ('Table 1', 'cx5', 4),
    ('A CX5', 'grey', 2, 1): ('Table 1', 'cx5_L2', 9),
    ('A CX5', 'grey', 2, 2): ('Table 1', 'cx5_L2 2°', 1),
    ('A CX6', 'grey', 1, None): ('Table 1', 'cx6', 6),
    ('B CX1', 'grey', 1, 1): ('Table 1', 'cx1_L1', 4),
    ('B CX1', 'grey', 1, 2): ('Table 1', 'cx1_L1 2°', 2),
    ('B CX1', 'grey', 2, 1): ('Table 1', 'cx1_L2', 23),
    ('B CX1', 'grey', 2, 2): ('Table 1', 'cx1_L2 2°', 4),
    ('B CX1', 'xray', None, None): ('Table 2', 'Bcx1', 5),
    ('B CX2', 'grey', 1, None): ('Table 1', 'cx2_L1', 5),
    ('B CX2', 'grey', 2, None): ('Table 1', 'cx2_L2', 5),
    ('B CX2', 'xray', 1, 1): ('Table 2', 'Bcx2_L1', 20),
    ('B CX2', 'xray', 1, 2): ('Table 2', 'Bcx2_L1 2nd', 23),
    ('B CX3', 'grey', None, None): ('Table 1', 'cx3', 33),
    ('B CX10', 'grey', 1, 1): ('Table 1', 'cx10_L1', 15),
    ('B CX10', 'grey', 1, 2): ('Table 1', 'cx10_L1 2°', 24),
    ('B CX10', 'grey', 2, None): ('Table 1', 'cx10_L2', 8),
    ('B CX10', 'xray', None, 1): ('Table 2', 'Bcx10_L1', 4),
    ('B CX10', 'xray', None, 2): ('Table 2', 'Bcx10_L2', 7),
    ('D CX3', 'grey', 2, None): ('Table 1', 'cx3_L2', 32),
    ('D CX6', 'grey', 1, None): ('Table 1', 'cx6_L1', 151),
    ('D CX6', 'grey', 2, None): ('Table 1', 'cx6_L2', 95),
    ('DCX7', 'grey', None, None): ('Table 1', 'cx7', 8),
    ('DCX7', 'xray', None, None): ('Table 2', 'Dcx7', 3),
}
# These traverses cross two compositional steps and the paper does not say which part
# it fitted. For the B CX1 X-ray scan the inner step gives 2.1-2.7 yr and the outer step
# 4.4 yr (Diffusor fits to 0-12, 0-10 and 10-23.2 um), against 5 yr in Table 2.
IOVINE_NO_SETUP = {('A CX5', 'grey', 1, None), ('B CX1', 'xray', None, None)}


def _blocks(rows_, i):
    """Split the distance row i of a transposed Iovine sheet into single traverses."""
    def label_row(name):
        for k in range(i - 1, max(i - 4, -1), -1):
            if rows_[k] and str(rows_[k][0]).strip().lower() == name:
                return rows_[k]
    line, prof = label_row('line'), label_row('profile')
    tag = lambda r, c: (int(r[c]) if r is not None and c < len(r) and number(r[c]) else None)
    dist = rows_[i]
    for j in range(i + 1, len(rows_)):
        values = rows_[j]
        if not values or values[0] is None:
            break
        label = str(values[0]).strip().lower()
        kind = {'gray value': 'grey', 'count rates': 'xray', 'bao': 'emp'}.get(label)
        if kind is None:
            continue
        cols = [c for c in range(1, min(len(dist), len(values))) if number(dist[c]) and number(values[c])]
        block, sign = [], 0
        for c in cols:
            if block:
                prev = block[-1]
                step = dist[c] - dist[prev]
                new_sign = (step > 0) - (step < 0)
                if (c != prev + 1 or tag(line, c) != tag(line, prev) or tag(prof, c) != tag(prof, prev)
                        or (sign and new_sign and new_sign != sign)):
                    yield kind, tag(line, block[0]), tag(prof, block[0]), j, block
                    block, sign, new_sign = [], 0, 0
                sign = new_sign or sign
            block.append(c)
        if block:
            yield kind, tag(line, block[0]), tag(prof, block[0]), j, block


def extract_iovine():
    """Iovine et al. (2017): BaO, greyscale and X-ray traverses across sanidine rims."""
    from openpyxl.utils import get_column_letter
    folder = SOURCE/'iovine2017'
    if not folder.exists():
        return
    for name in ('445_2017_1101_MOESM1_ESM.xlsx', '445_2017_1101_MOESM2_ESM.xlsx', '445_2017_1101_MOESM3_ESM.xlsx'):
        p = folder/name
        for sheet in sheets(p)[1:]:
            rr = [list(r) for r in rows(p, sheet)]
            seen = {}
            for i, r in enumerate(rr):
                if not (r and isinstance(r[0], str) and r[0].lower().startswith('distance')):
                    continue
                for kind, line, prof, j, block in _blocks(rr, i):
                    if len(block) < 6:
                        continue
                    ident = (sheet.strip(), kind, line, prof)
                    seen[ident] = seen.get(ident, 0) + 1
                    value = {'grey': 'Grey_value', 'xray': 'Ba_counts', 'emp': 'BaO_wt'}[kind]
                    points = [{'Distance_um': rr[i][c], value: rr[j][c],
                               'source_cell': f'{get_column_letter(c+1)}{j+1}'} for c in block]
                    label = (f"{sheet.strip()} {kind}" + (f" L{line}" if line else '')
                             + (f" P{prof}" if prof else '') + (f" #{seen[ident]}" if seen[ident] > 1 else ''))
                    published = IOVINE_PUBLISHED.get(ident) if seen[ident] == 1 else None
                    setup = None
                    if published and ident not in IOVINE_NO_SETUP:
                        setup = dict(T_C=930.0, coefficient='kfs_Ba_cherniak2002',
                                     published_years=published[2], published_table=published[0],
                                     published_label=published[1], member=sheet.strip()[0])
                    what = {'grey': 'Greyscale (accumulated BSE) swath', 'xray': 'Ba X-ray line scan (count rates)',
                            'emp': 'EMP BaO traverse'}[kind]
                    notes = (f'{what}, {name}, sheet {sheet!r}, distance in row {i+1} (headed "mm", but the '
                             f'values are micrometres), values in row {j+1}. ')
                    if published:
                        notes += (f'{published[0]} of the paper gives {published[2]} yr for {sheet.strip()[0]}-member '
                                  f'{published[1]} at 930 C (Cherniak 2002 Ba law, sharp initial step). ')
                    if ident in IOVINE_NO_SETUP:
                        notes += ('This traverse crosses two compositional steps and the paper does not say '
                                  'which part it fitted, so no setup is attached. ')
                    if kind == 'emp':
                        notes += ('The BaO fits of Table 1 used a part of the traverse that the paper does not '
                                  'tabulate, so no setup is attached. ')
                    spec = dict(distance_column='Distance_um', column_a=value, mode='A', distance_unit='um')
                    emit_table('iovine2017', label, p, sheet, points, mineral='kfeldspar', species='Ba',
                               value_column=value, spec=spec, setup=setup, notes=notes.strip())


def extract_sato():
    """Sato et al. (2022) Supplementary Table 8: EPMA line traverses of Zao pyroxenes."""
    p = SOURCE/'sato2022/1-s2.0-S0377027322002177-mmc1.xlsx'
    if not p.exists():
        return
    times = {}
    for i, r in enumerate(rows(p, 'Table9'), 1):
        if i > 2 and r[3] == 'EPMA' and number(r[8]):
            times[(r[0], r[2])] = dict(zone=r[4], MEs=r[5], T_C=r[6], buffer=r[7], residence_days=r[8],
                                       minimum_days=r[9], maximum_days=r[10], source_row=i)
    groups = {}
    for i, r in enumerate(rows(p, 'Table8'), 1):
        if i > 3 and isinstance(r[0], str) and number(r[4]) and number(r[29]):
            groups.setdefault((r[0], r[2], r[3]), []).append((i, r))
    for (sample, crystal, no), data in groups.items():
        wo = sorted(r[30] for _, r in data if number(r[30]))[len(data) // 2]
        mineral = 'opx' if wo < 10 else 'cpx'
        points = [dict(Distance_um=r[4], Mg_number=r[29], FeO_wt=r[8], MgO_wt=r[10], CaO_wt=r[12],
                       Wo=r[30], source_row=i) for i, r in data]
        modelled = times.get((sample, crystal))
        notes = ('EPMA line analysis, Supplementary Table 8. The Distance column carries no unit '
                 '(micrometres by the 1 um beam and 1-20 unit steps). ')
        if modelled:
            notes += (f"Supplementary Table 9 row {modelled['source_row']} gives an EPMA-based residence time of "
                      f"{modelled['residence_days']:g} d ({modelled['minimum_days']:g}-{modelled['maximum_days']:g}) "
                      f"for this crystal at {modelled['T_C']} C and {modelled['buffer']}, from the "
                      f"{modelled['zone']} ({modelled['MEs']}). The fitted part of the traverse and the step "
                      'position are not published, and fits to different parts of the traverse differ '
                      'by more than an order of magnitude, so no setup is attached.')
        else:
            notes += 'No EPMA-based time in Supplementary Table 9 for this crystal.'
        emit_table('sato2022', f'{sample} {crystal} No{no}', p, 'Table8', points, mineral=mineral,
                   species='Fe-Mg', value_column='Mg_number',
                   spec=dict(distance_column='Distance_um', column_a='FeO_wt', column_b='MgO_wt',
                             mode='A/(A+B)', oxide_a='FeO', oxide_b='MgO', distance_unit='um'),
                   notes=notes, extra=dict(source_timescale=modelled))


def extract_morgado():
    """Morgado et al. (2019) SM1: EPMA traverses across ilmenite-titanomagnetite pairs."""
    p = SOURCE/'morgado2019/410_2019_1596_MOESM1_ESM/SM1 - Fe-Ti oxide compositions.xlsx'
    if not p.exists():
        return
    for sheet in sheets(p):
        rr = rows(p, sheet)
        head = [str(v).strip().upper() if v is not None else '' for v in rr[0]]
        # column B holds the phase, headed PHASE or NUMBER
        iP, iD, iX = 1, head.index('RELDIST'), head.index('XTI')
        iTi, iFe = head.index('TIO2'), head.index('FEOT')
        points = []
        for n, r in enumerate(rr[1:], 2):
            if not (number(r[iD]) and number(r[iX])):
                continue
            phase = r[iP] if r[iP] in ('Ilm', 'Ti-mgt') else ''
            if phase == 'Ti-mgt' or (not phase and r[iX] < 0.35):
                points.append(dict(Distance_um=r[iD], X_Ti=r[iX], TiO2_wt=r[iTi], FeOt_wt=r[iFe],
                                   Phase=phase, source_row=n))
        emit_table('morgado2019', sheet, p, sheet, points, mineral='magnetite', species='Fe-Ti',
                   value_column='X_Ti',
                   spec=dict(distance_column='Distance_um', column_a='X_Ti', mode='A', distance_unit='um'),
                   notes=('Titanomagnetite rows of an ilmenite-titanomagnetite EPMA traverse (RELDIST, '
                          'micrometres, ilmenite rows omitted). Rows without a phase label are kept when '
                          'X_Ti < 0.35. The paper models these with the Aragon et al. (1984) point-defect '
                          'law and an interface held at equilibrium. The far boundary, the domain length '
                          'and the forward model are not published, so no setup is attached.'))


if __name__ == '__main__':
    extract()
