"""Extract the Kizimen orthopyroxene traverse shipped as a measured example.

Source: Ostorero et al. (2022) Communications Earth & Environment 3:290,
Supplementary Data 2, sheet "Opx profiles - compositions And" (microprobe
traverses) and Supplementary Data 4, sheet "Timescales opx Andesite" (their
modelled timescales). Both are archived at https://doi.org/10.5281/zenodo.7307563.

The crystal is K9_L10C4 (andesite K9, 355 um sieve fraction), a single reverse
zone modelled by Ostorero et al. at 2.32 years (+7.16 / -1.75). Rows are kept
from the first point the authors flagged as good quality (green in the
spreadsheet) to the end of the traverse; nothing is smoothed or re-scaled.

Usage:
    python scripts/extract_kizimen.py "path/to/ostorero_kizimen_data"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

CRYSTAL = "K9_L10C4"
GOOD_FILL = "FF92D050"          # the green row colour used in the spreadsheet
OUT = Path(__file__).resolve().parent.parent / "examples" / "opx_kizimen_ostorero2022.csv"


def main(folder: Path) -> None:
    wb = openpyxl.load_workbook(folder / "Supplementary Data 2.xlsx", data_only=True)
    ws = wb["Opx profiles - compositions And"]
    header = [ws.cell(2, c).value for c in range(1, 19)]
    # row 3 holds the authors' mean analytical standard deviation per oxide (wt%)
    sd = {header[c - 1]: ws.cell(3, c).value for c in range(3, 15)}
    rows, started = [], False
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 2).value != CRYSTAL:
            continue
        fill = ws.cell(r, 2).fill.fgColor.rgb if ws.cell(r, 2).fill.fgColor else None
        started = started or fill == GOOD_FILL
        if not started:
            continue
        v = {h: ws.cell(r, c).value for c, h in enumerate(header, start=1)}
        rows.append({
            "Distance_from_rim_um": v["Distance from the rim (µm)"],
            "FeO_wt": v["FeO"], "MgO_wt": v["MgO"],
            "FeO_err": sd["FeO"], "MgO_err": sd["MgO"],
            "CaO_wt": v["CaO"], "Al2O3_wt": v["Al2O3"], "SiO2_wt": v["SiO2"],
            "MnO_wt": v["MnO"], "Total_wt": v["Total"],
            "En_percent": v["En (%)"], "Mg_number": v["Mg#"],
        })
    df = pd.DataFrame(rows)
    df.to_csv(OUT, index=False, float_format="%.5g")
    print(f"wrote {len(df)} rows for {CRYSTAL} to {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
