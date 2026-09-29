"""Profiles read from images: readers, geometry, cleaning, colour legends, workbook, GUI."""
import os

import numpy as np
import pandas as pd
import pytest
from scipy.special import erf

from diffusor.dataio import build_profile, calibrate
from diffusor.dataio.image_profiles import (EXCLUDED, KEPT, LINE_DROPPED, OUTLIER, THRESHOLD,
                                            CleaningSettings, ExtractionSettings, _clip,
                                            composition_table, extract_profiles,
                                            is_extraction_workbook, read_extraction,
                                            table_from_extraction, write_workbook)
from diffusor.dataio.images import ColorScale, LoadedImage, load_image, value_map

H, W = 300, 360
START, END = (60.0, 260.0), (300.0, 20.0)       # guideline running up and to the right


def _edge(width=8.0, noise=2.0, seed=1):
    """Brighter on the right of the guideline (looking from START to END)."""
    yy, xx = np.mgrid[0:H, 0:W].astype(float)
    d = np.array(END) - np.array(START)
    d /= np.hypot(*d)
    right = np.array([-d[1], d[0]])
    signed = (xx - START[0]) * right[0] + (yy - START[1]) * right[1]
    img = 100 + 80 * 0.5 * (1 + erf(signed / width))
    img += np.random.default_rng(seed).normal(0, noise, img.shape)
    return img, signed


def _true(offset_px, width=8.0):
    # offset is negative on the start (right) side, where signed distance is positive
    return 100 + 80 * 0.5 * (1 + erf(-offset_px / width))


def _grey(arr, name="synthetic.png"):
    return LoadedImage(np.asarray(arr, float), name, "test", ["value"], (0.0, 255.0))


# --------------------------------------------------------------------- geometry
def test_first_value_is_right_of_the_guideline_as_in_nidis():
    img, _ = _edge(noise=0.0)
    ex = extract_profiles(_grey(img), ExtractionSettings([START, END],
                                                         cleaning=CleaningSettings(clip_method="none")))
    mean = ex.statistics("clean")["Mean"].to_numpy()
    assert mean[0] > 170 and mean[-1] < 110            # starts on the bright right side
    assert ex.n_samples == 101 and ex.n_lines == round(np.hypot(240, 240)) + 1
    assert np.allclose(ex.offset_px[[0, -1]], [-50, 50])
    assert np.max(np.abs(mean - _true(ex.offset_px))) < 0.5
    flipped = extract_profiles(_grey(img), ExtractionSettings([START, END], flip=True,
                                                              cleaning=CleaningSettings(clip_method="none")))
    assert flipped.statistics()["Mean"].iloc[0] < 110


def test_lines_are_perpendicular_to_a_curved_guideline():
    yy, xx = np.mgrid[0:H, 0:W].astype(float)
    r = np.hypot(xx - 180, yy - 150)
    img = 100 + 80 * 0.5 * (1 + erf((r - 90) / 6.0))       # bright outside a circle
    arc = [(180 + 90 * np.cos(t), 150 + 90 * np.sin(t)) for t in np.linspace(-0.8, 0.8, 25)]
    ex = extract_profiles(_grey(img), ExtractionSettings(arc, 30, 30,
                                                         cleaning=CleaningSettings(clip_method="none")))
    st = ex.statistics()
    # lines cross the circle at right angles, so every line sees the same profile
    assert st["SD"].max() < 1.5
    # the arc runs down the right side of the circle on screen, so looking along it the
    # right-hand side is the circle's centre: lines start inside, in the dark
    assert st["Mean"].iloc[0] < 110 and st["Mean"].iloc[-1] > 170


def test_samples_off_the_image_are_not_used():
    img, _ = _edge(noise=0.0)
    ex = extract_profiles(_grey(img), ExtractionSettings([(5, 150), (5, 100)], 20, 20))
    assert np.isnan(ex.raw).any()
    assert ex.statistics("raw")["N"].min() == 0
    assert any("outside the image" in n for n in ex.notes)


# --------------------------------------------------------------------- cleaning
def test_cracks_and_inclusions_are_removed():
    img, _ = _edge()
    yy, xx = np.mgrid[0:H, 0:W]
    img[np.abs(yy - 0.4 * xx - 90) < 1.5] = 4           # a dark crack across the lines
    img[(xx - 200) ** 2 + (yy - 110) ** 2 < 49] = 252   # a bright inclusion
    raw_only = extract_profiles(_grey(img), ExtractionSettings(
        [START, END], cleaning=CleaningSettings(clip_method="none", max_rejected_fraction=None)))
    cleaned = extract_profiles(_grey(img), ExtractionSettings(
        [START, END], cleaning=CleaningSettings(low=40, high=235, grow_px=1)))
    truth = _true(cleaned.offset_px)
    err_raw = np.max(np.abs(raw_only.statistics()["Mean"] - truth))
    err_clean = np.max(np.abs(cleaned.statistics()["Mean"] - truth))
    assert err_raw > 2.0 and err_clean < 0.6
    assert np.sum(cleaned.reason == THRESHOLD) > 0
    # the raw statistics keep everything that is inside the image
    inside = np.sum(cleaned.reason != 1, axis=1)
    assert np.array_equal(cleaned.statistics("raw")["N"].to_numpy(), inside)


def test_nidis_one_sigma_test_removes_a_third_of_clean_data_and_mad_almost_none():
    v = np.random.default_rng(0).normal(0, 1, (40, 500))
    ok = np.ones(v.shape, bool)
    assert 0.29 < _clip(v, ok, "nidis", 1.0, 1).mean() < 0.35
    assert _clip(v, ok, "mad", 3.0, 5).mean() < 0.01


def test_exclusion_area_rejects_what_it_covers():
    img, _ = _edge()
    poly = [(150, 120), (190, 120), (190, 160), (150, 160)]
    ex = extract_profiles(_grey(img), ExtractionSettings(
        [START, END], cleaning=CleaningSettings(exclusions=[poly], clip_method="none")))
    hit = ex.reason == EXCLUDED
    assert hit.any()
    assert np.all((ex.xs[hit] > 148) & (ex.xs[hit] < 192) & (ex.ys[hit] > 118) & (ex.ys[hit] < 162))


def test_a_lamella_along_the_lines_drops_those_lines():
    img, _ = _edge(noise=1.0)
    yy, xx = np.mgrid[0:H, 0:W].astype(float)
    # a dark band parallel to the profile lines (perpendicular to the guideline)
    img[np.abs((xx - 180) - (yy - 140)) < 3] = 10
    ex = extract_profiles(_grey(img), ExtractionSettings(
        [START, END], cleaning=CleaningSettings(low=40, max_rejected_fraction=0.5)))
    assert ex.dropped_lines.sum() >= 3
    assert np.any(ex.reason == LINE_DROPPED)


# --------------------------------------------------------------------- readers
def test_readers_cover_common_formats(tmp_path):
    from PIL import Image
    grey8 = (np.arange(40 * 30).reshape(30, 40) % 256).astype(np.uint8)
    Image.fromarray(grey8).save(tmp_path / "a.png")
    Image.fromarray(grey8).save(tmp_path / "a.bmp")
    g16 = (np.arange(40 * 30).reshape(30, 40) * 50).astype(np.uint16)
    Image.fromarray(g16).save(tmp_path / "b.tif")
    rgb = np.dstack([grey8, 255 - grey8, np.full_like(grey8, 7)])
    Image.fromarray(rgb).save(tmp_path / "c.jpg", quality=100)
    Image.fromarray(rgb).save(tmp_path / "c.png")
    f32 = np.linspace(0, 1, 1200, dtype=np.float32).reshape(30, 40)
    Image.fromarray(f32).save(tmp_path / "d.tif")
    np.save(tmp_path / "e.npy", f32)

    assert np.array_equal(load_image(tmp_path / "a.png").data, grey8)
    assert np.array_equal(load_image(tmp_path / "a.bmp").data, grey8)
    b = load_image(tmp_path / "b.tif")
    assert b.value_range == (0.0, 65535.0) and np.array_equal(b.data, g16)
    c = load_image(tmp_path / "c.png")
    assert c.is_rgb
    lum, _, _ = value_map(c, "luminance")
    assert np.allclose(lum, 0.299 * grey8 + 0.587 * (255 - grey8.astype(float)) + 0.114 * 7)
    assert load_image(tmp_path / "c.jpg").n_channels == 3
    assert np.allclose(load_image(tmp_path / "d.tif").data, f32)
    assert np.allclose(load_image(tmp_path / "e.npy").data, f32)


def test_multipage_tiff_becomes_channels(tmp_path):
    from PIL import Image
    pages = [Image.fromarray(np.full((20, 25), v, np.uint16)) for v in (100, 2000, 30000)]
    pages[0].save(tmp_path / "maps.tif", save_all=True, append_images=pages[1:])
    im = load_image(tmp_path / "maps.tif")
    assert im.n_channels == 3 and im.channel_names[0] == "page 1"
    vals, _, _ = value_map(im, "channel", 2)
    assert np.all(vals == 30000)


def test_text_grid_envi_and_raw(tmp_path):
    grid = np.arange(12 * 9, dtype=float).reshape(9, 12)
    with open(tmp_path / "map.txt", "w") as fh:
        fh.write("Mg Ka counts, JEOL export\n")
        for r in grid:
            fh.write("\t".join(f"{v:g}" for v in r) + "\n")
    assert np.array_equal(load_image(tmp_path / "map.txt").data, grid)
    pd.DataFrame(grid).to_csv(tmp_path / "map.csv", index=False, header=False)
    assert np.array_equal(load_image(tmp_path / "map.csv").data, grid)

    data = (np.arange(16 * 10) * 7).astype("<u2").reshape(10, 16)
    data.tofile(tmp_path / "bse.raw")
    (tmp_path / "bse.hdr").write_text(
        "ENVI\nsamples = 16\nlines = 10\nbands = 1\nheader offset = 0\n"
        "data type = 12\ninterleave = bsq\nbyte order = 0\n")
    assert np.array_equal(load_image(tmp_path / "bse.hdr").data, data)
    assert np.array_equal(load_image(tmp_path / "bse.raw").data, data)

    (tmp_path / "unknown.bhd").write_bytes(b"\x00" * 8 + data.astype(">u2").tobytes())
    with pytest.raises(ValueError, match="raw"):
        load_image(tmp_path / "unknown.bhd")
    raw = load_image(tmp_path / "unknown.bhd", raw=dict(width=16, height=10, dtype="uint16",
                                                         byte_order="big", offset=8))
    assert np.array_equal(raw.data, data)


def test_pixel_size_is_read_from_metadata(tmp_path):
    from PIL import Image
    from PIL.TiffImagePlugin import ImageFileDirectory_v2
    arr = np.zeros((10, 10), np.uint8)
    ifd = ImageFileDirectory_v2()
    ifd[270] = "ImageJ=1.54f\nunit=micron\n"
    ifd[282] = 4.0                          # 4 pixels per micron
    ifd[283] = 4.0
    Image.fromarray(arr).save(tmp_path / "ij.tif", tiffinfo=ifd)
    im = load_image(tmp_path / "ij.tif")
    assert im.pixel_size_um == pytest.approx(0.25) and "ImageJ" in im.pixel_size_source

    ifd = ImageFileDirectory_v2()
    ifd[34682] = "[Scan]\nPixelWidth=5.2e-008\nPixelHeight=5.2e-008\n"
    ifd.tagtype[34682] = 2
    Image.fromarray(arr).save(tmp_path / "fei.tif", tiffinfo=ifd)
    assert load_image(tmp_path / "fei.tif").pixel_size_um == pytest.approx(0.052)

    Image.fromarray(arr).save(tmp_path / "jeol.tif")
    (tmp_path / "jeol.txt").write_text("$CM_FULL_SIZE 1280 960\n$$SM_MICRON_BAR 200\n"
                                       "$$SM_MICRON_MARKER 50um\n")
    assert load_image(tmp_path / "jeol.tif").pixel_size_um == pytest.approx(0.25)


# --------------------------------------------------------------------- colour scale
def _rainbow_map(vmin=0.0, vmax=40.0):
    from matplotlib import colormaps
    img, _ = _edge(noise=0.0)
    values = vmin + (img - 100) / 80 * (vmax - vmin)            # 0-40 wt%
    rgb = colormaps["jet"]((values - vmin) / (vmax - vmin))[:, :, :3]
    return values, rgb


def test_named_colour_map_is_inverted_and_off_scale_pixels_flagged():
    values, rgb = _rainbow_map()
    rgb[100:103, :, :] = 0.0                                     # black crack
    rgb[200:204, :, :] = 1.0                                     # white label
    cs = ColorScale.from_colormap("jet", 0.0, 40.0, units="wt% MgO")
    got, dist = cs.to_values(rgb)
    good = np.isfinite(got)
    assert not good[100:103].any() and not good[200:204].any()
    assert np.nanmax(np.abs(got - values)[good]) < 0.25          # well under 1% of the range


def test_legend_drawn_in_the_image_gives_values(tmp_path):
    from matplotlib import colormaps
    from PIL import Image
    values, rgb = _rainbow_map()
    bar = colormaps["jet"](np.linspace(0, 1, 200))[:, :3]
    rgb = rgb.copy()
    rgb[5:17, 10:210, :] = bar[None, :, :]                       # legend top left, low end left
    Image.fromarray((rgb * 255).round().astype(np.uint8)).save(tmp_path / "map.png")
    im = load_image(tmp_path / "map.png")
    cs = ColorScale.from_legend(im, (10, 11), (209, 11), 0.0, 40.0, units="wt%")
    ex = extract_profiles(im, ExtractionSettings([START, END], 40, 40, value_mode="colour_scale",
                                                 cleaning=CleaningSettings(clip_method="none")),
                          color_scale=cs)
    truth = (_true(ex.offset_px) - 100) / 80 * 40
    assert np.nanmax(np.abs(ex.statistics()["Mean"] - truth)) < 0.5
    assert "wt%" in ex.value_label


# --------------------------------------------------------------------- workbook
def test_workbook_round_trip_and_calibrated_profile(tmp_path):
    img, _ = _edge()
    s = ExtractionSettings([START, END], 30, 30, pixel_size_um=0.1,
                           cleaning=CleaningSettings(low=20, high=240))
    ex = extract_profiles(_grey(img), s)
    path = write_workbook(tmp_path / "p.xlsx", ex, overlay_png=None)
    from openpyxl import load_workbook
    names = load_workbook(path, read_only=True).sheetnames
    assert names[0] == "Profile"
    for sheet in ("Summary", "Raw lines", "Clean lines", "Rejection codes", "Geometry", "Settings"):
        assert sheet in names
    assert is_extraction_workbook(path) and not is_extraction_workbook(tmp_path / "missing.xlsx")
    tb = read_extraction(path)
    assert tb.pixel_size_um == pytest.approx(0.1)
    assert tb.settings.guideline == [START, END] and tb.settings.cleaning.low == 20
    lines = pd.read_excel(path, sheet_name="Raw lines")
    assert lines.shape[1] == ex.n_lines + 3

    # grey 100 -> X_Fe 0.30, grey 180 -> X_Fe 0.10
    cal = calibrate([100, 140, 180], [0.30, 0.20, 0.10])
    df = composition_table(tb, tb.pixel_size_um, "Clean_Mean", "Clean_SE", cal, "XFe")
    assert df["Distance_um"].iloc[-1] == pytest.approx(6.0)
    expected = 0.30 - (_true(ex.offset_px) - 100) / 80 * 0.20
    assert np.max(np.abs(df["XFe"] - expected)) < 0.003
    assert np.all(df["XFe_err"] > 0)
    from diffusor.dataio import ProfileSpec
    prof = build_profile(df, ProfileSpec("Distance_um", "XFe", sigma_a_column="XFe_err", mode="A"))
    assert len(prof) == ex.n_samples

    # the same table without a file, and a missing scale is refused
    assert list(table_from_extraction(ex).profile.columns) == list(tb.profile.columns)
    with pytest.raises(ValueError, match="pixel size"):
        composition_table(tb, None, calibration=cal)


# --------------------------------------------------------------------- interface
@pytest.fixture(scope="module")
def app():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_extractor_window_draws_saves_and_loads_into_diffusor(app, tmp_path, monkeypatch):
    from PIL import Image
    from diffusor.gui.image_calibration import ImageCalibrationDialog
    from diffusor.gui.main_window import MainWindow
    from PySide6.QtWidgets import QDialog

    img, _ = _edge()
    Image.fromarray(np.clip(img, 0, 255).astype(np.uint8)).save(tmp_path / "bse.png")
    w = MainWindow()
    w.show_image_extractor()
    dlg = w.image_extractor
    dlg.open_image(str(tmp_path / "bse.png"))
    dlg.sp_px.setValue(0.05)
    dlg.guideline = [START, END]
    dlg.recompute()
    assert dlg.extraction is not None and dlg.extraction.n_lines > 300
    out = dlg.save_workbook(str(tmp_path / "bse_profile1.xlsx"))
    assert out and is_extraction_workbook(out)

    def accept(self):
        self.cmb_map.setCurrentIndex(0)
        self.sp_g1.setValue(100.0)
        self.sp_c1.setValue(0.3)
        self.sp_g2.setValue(180.0)
        self.sp_c2.setValue(0.1)
        self._accept()
        return QDialog.Accepted if self.frame is not None else QDialog.Rejected

    monkeypatch.setattr(ImageCalibrationDialog, "exec", accept)
    w._load(tmp_path / "bse_profile1.xlsx", None)         # the File > Load path
    assert w.profile is not None and len(w.profile) == dlg.extraction.n_samples
    assert w.profile.C.min() > 0.08 and w.profile.C.max() < 0.32
    w.profile = None
    dlg.use_in_diffusor()                                   # the extractor's button
    assert w.profile is not None
    dlg.close()
    w.close()


def test_extractor_handles_a_colour_map_with_a_drawn_legend(app, tmp_path):
    from matplotlib import colormaps
    from PIL import Image
    from diffusor.gui.image_extractor import ImageExtractorDialog
    _, rgb = _rainbow_map()
    rgb = rgb.copy()
    rgb[5:17, 10:210, :] = colormaps["jet"](np.linspace(0, 1, 200))[None, :, :3]
    Image.fromarray((rgb * 255).round().astype(np.uint8)).save(tmp_path / "mg.png")
    dlg = ImageExtractorDialog()
    dlg.open_image(str(tmp_path / "mg.png"))
    dlg.cmb_mode.setCurrentIndex(dlg.cmb_mode.findData("colour_scale"))
    dlg.sp_v1.setValue(40.0)
    dlg.legend_pts = [(10, 11), (209, 11)]
    dlg._legend_changed()
    dlg.guideline = [START, END]
    dlg.recompute()
    assert dlg.extraction is not None and dlg.extraction.color_scale is not None
    mean = dlg.extraction.statistics()["Mean"]
    assert mean.iloc[0] > 35 and mean.iloc[-1] < 5
    dlg.close()


def test_extractor_uses_a_legend_saved_as_a_separate_image(app, tmp_path):
    from types import SimpleNamespace
    from matplotlib import colormaps
    from PIL import Image
    from PySide6.QtWidgets import QDialogButtonBox
    from diffusor.gui.image_extractor import ImageExtractorDialog, LegendPickerDialog
    _, rgb = _rainbow_map()                                     # the map carries no legend
    Image.fromarray((rgb * 255).round().astype(np.uint8)).save(tmp_path / "mg.png")
    legend = np.ones((60, 240, 3))                              # white margin round the bar
    legend[20:40, 20:220] = colormaps["jet"](np.linspace(0, 1, 200))[None, :, :3]
    Image.fromarray((legend * 255).round().astype(np.uint8)).save(tmp_path / "legend.png")

    from diffusor.dataio.images import load_image
    leg = load_image(tmp_path / "legend.png")
    picker = LegendPickerDialog(leg)
    assert not picker.bb.button(QDialogButtonBox.Ok).isEnabled()
    for x, y in ((20, 30), (219, 30)):
        picker._on_click(SimpleNamespace(inaxes=picker.ax, xdata=x, ydata=y, button=1))
    assert picker.points == [(20.0, 30.0), (219.0, 30.0)]
    assert picker.bb.button(QDialogButtonBox.Ok).isEnabled()
    picker.close()

    dlg = ImageExtractorDialog()
    dlg.open_image(str(tmp_path / "mg.png"))
    dlg.cmb_mode.setCurrentIndex(dlg.cmb_mode.findData("colour_scale"))
    dlg.sp_v1.setValue(40.0)
    dlg.set_legend_image(leg, picker.points)
    assert dlg.cmb_legend.currentData() == "file"
    assert not dlg.tool_buttons["legend"].isEnabled()       # nothing to draw on the map itself
    dlg.guideline = [START, END]
    dlg.recompute()
    cs = dlg.extraction.color_scale
    assert "legend.png" in cs.source
    truth = (_true(dlg.extraction.offset_px) - 100) / 80 * 40
    assert np.nanmax(np.abs(dlg.extraction.statistics()["Mean"] - truth)) < 0.5
    dlg.close()


def test_image_windows_fit_the_screen(app):
    from PySide6.QtGui import QGuiApplication
    from diffusor.gui.image_calibration import ImageCalibrationDialog
    from diffusor.gui.image_extractor import ImageExtractorDialog
    avail = QGuiApplication.primaryScreen().availableGeometry()
    img, _ = _edge()
    ex = extract_profiles(_grey(img), ExtractionSettings([START, END]))
    for dlg in (ImageExtractorDialog(), ImageCalibrationDialog(table_from_extraction(ex))):
        assert dlg.width() <= avail.width() and dlg.height() <= avail.height()
        # small screens: the content scrolls rather than pushing the buttons off screen
        assert dlg.minimumSizeHint().height() < 450
        dlg.close()
