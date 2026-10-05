import numpy as np
import pytest
from diffusor.coefficients import Conditions, get
from diffusor.fitting import DiffusionModel
from diffusor.solvers.initial import InitialCondition
from diffusor.fitting.joint import ProfileConstraint, fit_joint_time


def constraint(name, time, scale=1.):
    model = DiffusionModel(get('qz_Ti_cherniak2007'), Conditions(1373.15, axis='c'),
        InitialCondition('step', {'x0':0., 'C_left':scale, 'C_right':0.}))
    x=np.linspace(-80,80,61)
    return ProfileConstraint(name,model,x,model.profile(time,x),np.full(x.size,.015*scale))


def test_shared_duration_recovers_synthetic_time_and_is_invariant_to_units():
    a,b=constraint('a',1e8),constraint('b',1e8,1000)
    r=fit_joint_time([a,b],t_min=1e6,t_max=1e10,scan_points=30)
    assert r.t_seconds==pytest.approx(1e8,rel=1e-5)
    assert r.dof==121
    assert r.chi2==pytest.approx(sum(f.stats.chi2 for f in r.profiles),abs=1e-9)
    b2=constraint('b',2e8,1000)
    r2=fit_joint_time([a,b2],t_min=1e6,t_max=1e10,scan_points=30)
    r3=fit_joint_time([a,constraint('b',2e8)],t_min=1e6,t_max=1e10,scan_points=30)
    assert r2.t_seconds==pytest.approx(r3.t_seconds,rel=1e-7)
    assert 1e8<r2.t_seconds<2e8
    assert r2.reduced_chi2>1


def test_joint_uncertainty_validation_and_cancellation():
    a=constraint('a',1e8)
    for sigma in (None,0.,-1.,float('nan')):
        with pytest.raises(ValueError):
            ProfileConstraint('bad',a.model,a.x,a.concentration,sigma)
    with pytest.raises(ValueError,match='unique'):
        fit_joint_time([a,a])
    with pytest.raises(InterruptedError):
        fit_joint_time([a,constraint('b',1e8)],progress=lambda _:True)


def test_joint_export_contains_each_profile_and_global_degrees_of_freedom(tmp_path):
    from diffusor.dataio.joint import save_joint_results
    from openpyxl import load_workbook
    import json
    result=fit_joint_time([constraint('a',1e8),constraint('b',1e8)],t_min=1e7,t_max=1e9,scan_points=15)
    files=save_joint_results(tmp_path,result)
    data=json.loads((tmp_path/'joint_results.json').read_text())
    assert data['dof']==121
    assert len(data['profiles'])==2
    wb=load_workbook(files['workbook'],data_only=True)
    assert wb['Profiles'].max_row==3
    assert (tmp_path/'profile_001/diffusor_run_methods.txt').exists()
    assert (tmp_path/'joint_profiles.svg').exists()


def test_joint_dialog_keeps_independent_snapshots_and_runs_worker():
    import os,time
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    from types import SimpleNamespace
    from PySide6.QtWidgets import QApplication,QWidget
    from diffusor.gui.joint_study import JointStudyDialog
    app=QApplication.instance() or QApplication([])
    class Owner(QWidget):
        def __init__(self):
            super().__init__()
            self.c=constraint('a',1e8)
            self.profile=SimpleNamespace(x=self.c.x,C=self.c.concentration,sigma=self.c.sigma)
        def _ready(self): return True
        def _fit_profile(self): return self.profile
        def _checked_keys(self): return ['selected']
        def _model(self,key): return self.c.model
    owner=Owner();dialog=JointStudyDialog(owner)
    dialog.add_current()
    dialog.add_current()
    assert len(dialog.profiles)==2
    owner.c.model.initial.params['C_left']=2
    assert dialog.profiles[0].model.initial.params['C_left']==1
    dialog.run_fit()
    deadline=time.monotonic()+20
    while dialog.thread.isRunning() and time.monotonic()<deadline:
        app.processEvents()
        time.sleep(.01)
    if dialog.thread.isRunning():
        dialog.abort();dialog.thread.quit();dialog.thread.wait(3000)
        pytest.fail('Joint worker did not finish')
    app.processEvents()
    assert dialog.result is not None
    assert dialog.result.t_seconds==pytest.approx(1e8,rel=1e-5)
    assert dialog.export.isEnabled()
    dialog.close();owner.close()
