"""Regression checks for audited equations, directions, and scientific reporting."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np
import pytest
from openpyxl import load_workbook

from diffusor.coefficients import Conditions, get, list_coefficients
from diffusor.coefficients.magnetite import SIEVWRIGHT_TABLE5
from diffusor.fitting.model import DiffusionModel
from diffusor.solvers.initial import InitialCondition
from diffusor.solvers.history import ThermalHistory


@pytest.mark.parametrize("regime,expected", [("tamed", -16.856), ("ped", -17.279)])
def test_olivine_equations_follow_the_erratum(regime, expected):
    # Eqs 27/28 (pp. 424-425) as corrected by the erratum (PCM 34:597-598), whose
    # composition term is 3 (XFe - 0.1): at Fo90 it vanishes. Evaluated at 1100 C,
    # XFe = 0.1, 1 bar, fO2 = 1e-7 Pa, tabulated by hand from the erratum. The printed
    # 3 XFe would give 0.3 log units more, far outside the tolerance.
    c = get("ol_FeMg_dohmen_chakraborty2007_" + regime)
    cond = Conditions(1373.15, X={"XFe": .1}, log_fo2_bar=-12, axis="c")
    assert c.log10_D(cond) == pytest.approx(expected, abs=.015)
    richer = c.log10_D(cond.replace(X={"XFe": .2}))
    assert richer - c.log10_D(cond) == pytest.approx(0.3)


def test_ni_matches_petry_experiment_ni10():
    # Petry Table 2: 1005 C, log fO2=-5.96 Pa, Fo90, log D=-17.12.
    c = get("ol_Ni_petry2004")
    assert c.log10_D(Conditions(1278.15, log_fo2_bar=-10.96, X={"XFe": .1}, axis="c")) == pytest.approx(-17.12, abs=.16)


def test_ca_redox_reference_and_axis_projection():
    c = get("ol_Ca_coogan2005")
    base = Conditions(1473.15, log_fo2_bar=-12, axis="c")
    assert c.log10_D(base) == pytest.approx(-17.36, abs=.01)
    assert c.D(base.replace(log_fo2_bar=-10)) / c.D(base) == pytest.approx(10**.62)
    projected = c.D(base.replace(axis=None, angles_deg=(45, 90, 45)))
    assert projected == pytest.approx(.5*(c.D(base)+c.D(base.replace(axis="a"))))


def test_be_anisotropy_is_temperature_dependent_and_orientation_mandatory():
    c = get("ol_Be_jollands2016")
    with pytest.raises(ValueError, match="orientation"):
        c.D(Conditions(1500))
    def ratio(T):
        return c.D(Conditions(T, axis="c"))/c.D(Conditions(T, axis="a"))
    assert ratio(1223.15) > 8 * ratio(1748.15)
    assert c.log10_D(Conditions(1473.15, axis="c")) == pytest.approx(-13.89, abs=.01)


def test_single_direction_law_cannot_silently_ignore_orientation():
    c = get("rt_Zr_cherniak2007_c")
    for cond in (Conditions(1300, axis="a"), Conditions(1300, angles_deg=(45, 90, 45))):
        with pytest.raises(ValueError, match="calibrated only"):
            c.D(cond)


def test_all_magnetite_rows_available_at_measured_temperature_only():
    for sp in SIEVWRIGHT_TABLE5:
        fixed = [c for c in list_coefficients("magnetite", sp)
                 if c.citation == "sievwright2020" and c.fixed_temperature_K is not None]
        assert len(fixed) == 1
        c = fixed[0]
        assert np.isfinite(c.D(Conditions(1423.15, log_fo2_bar=-8)))
        with pytest.raises(ValueError, match="no temperature dependence"):
            c.D(Conditions(1323.15, log_fo2_bar=-8))
    warnings = get("mt_Ti_sievwright2020").check_conditions(Conditions(1323.15, log_fo2_bar=-8))
    assert any("HYPOTHESIS" in w for w in warnings)
    assert any("outside" in w for w in warnings)


def test_fo_percent_maps_to_host_xfe_and_preserves_complementary_profile():
    c = get("ol_FeMg_dohmen_chakraborty2007_tamed")
    cond = Conditions(1423.15, log_fo2_bar=-12, X={"XFe": .15}, axis="c")
    x = np.linspace(0, 100, 101)
    fe = DiffusionModel(c, cond, InitialCondition("step", {"x0": 50, "C_left": .1, "C_right": .2}),
                        comp_key="XFe", x_grid=x)
    fo = DiffusionModel(c, cond, InitialCondition("step", {"x0": 50, "C_left": 90, "C_right": 80}),
                        comp_key="XFe", comp_scale=-.01, comp_offset=1., x_grid=x)
    assert np.allclose(fe.profile(1e6, x), 1-fo.profile(1e6, x)/100, atol=1e-10)


def test_isothermal_history_temperature_overrides_nominal_conditions():
    c = get("ol_P_watson2015")
    m = DiffusionModel(c, Conditions(1000), InitialCondition("step", {"x0": 0, "C_left": 1, "C_right": 0}),
                       history=ThermalHistory.isothermal(1100, 10))
    assert m._effective_Dt(10) == pytest.approx(c.D(Conditions(1100))*1e12*10)


def test_coefficient_provenance_survives_excel_and_json(tmp_path):
    from diffusor.fitting.fit import fit_time
    from diffusor.dataio.export import save_results, result_dict
    c = get("qz_Ti_cherniak2007")
    model = DiffusionModel(c, Conditions(1273.15, axis="c"),
                           InitialCondition("step", {"x0": 0, "C_left": 50, "C_right": 10}))
    x = np.linspace(-40, 40, 41)
    y = model.profile(1e8, x)
    result = fit_time(model, x, y, t_min=1e6, t_max=1e10)
    files = save_results(tmp_path, result)
    wb = load_workbook(files["workbook"], data_only=True)
    assert wb.sheetnames == ["Results", "Profile", "Coefficient", "Methods"]
    assert wb["Results"]["B5"].value == pytest.approx(result.t_seconds)
    assert len(wb["Results"]._charts) == 2
    assert wb["Profile"].freeze_panes == "A2"
    assert wb["Profile"].max_row == 42
    assert wb["Profile"]["B2"].value == pytest.approx(y[0])
    metadata = result_dict(result)["coefficient"]
    assert metadata["kind"] == "chemical"
    assert metadata["validation_level"] == "source_transcription_checked"
    assert metadata["allowed_axes"] == ["c"]


def test_new_minerals_are_available_in_gui_and_trace_does_not_drive_host_composition():
    from PySide6.QtWidgets import QApplication
    from diffusor.gui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    w = MainWindow()
    for key, sp in (("quartz", "Ti"), ("rutile", "Zr"), ("titanite", "Sr")):
        w.cmb_mineral.setCurrentIndex(w.cmb_mineral.findData(key))
        assert w.cmb_species.findText(sp) >= 0
        assert w.lst_coef.count() > 0
    w.cmb_mineral.setCurrentIndex(w.cmb_mineral.findData("olivine"))
    w.cmb_species.setCurrentText("Ni")
    assert w._model("ol_Ni_petry2004").comp_key is None
    w.close()


@pytest.mark.parametrize('geometry',[0,1,2])
def test_total_flux_and_exact_control_volume_mass_are_conserved(geometry):
    from diffusor.solvers.numerical import solve_1d
    x=np.linspace(0,10,35)
    C0=.3+.1*np.cos(x)
    an=.2+.05*x
    r=solve_1d(x,C0,lambda C,T: .4+np.asarray(C)*.2,10.,m=geometry,
               T_K=1200,an_profile=an,theta_activity=2.)
    assert r.mass_final == pytest.approx(r.mass_initial,rel=2e-12)
    assert np.min(r.C_final)>0


def test_activity_equilibrium_is_preserved_at_closed_boundaries():
    from diffusor.solvers.numerical import solve_1d
    x=np.linspace(0,10,50)
    an=.1+.07*x
    q=np.diff(an) # theta/2 = 1
    C0=np.r_[1.,np.cumprod((1+q)/(1-q))]
    r=solve_1d(x,C0,lambda C,T:1.,25.,T_K=1200,an_profile=an,theta_activity=2.)
    assert np.allclose(r.C_final,C0,rtol=1e-11,atol=1e-11)


def test_accessory_unit_conversion_and_zircon_directional_contrast():
    from diffusor.constants import R_GAS
    ap=get('ap_Sr_cherniak1993')
    # Original units: cm2/s, cal/mol. One cm2/s = 1e-4 m2/s.
    expected=2.7e-3*1e-4*np.exp(-65000*4.184/(R_GAS*1273.15))
    assert ap.D(Conditions(1273.15))==pytest.approx(expected,rel=1e-12)
    fast=get('zrn_Ti_bloch2022_c')
    slow=get('zrn_Ti_cherniak2007_perp_c')
    ratio=fast.D(Conditions(1723.15,axis='c'))/slow.D(Conditions(1723.15,axis='a'))
    assert 1e4<ratio<1e5
    with pytest.raises(ValueError,match='orientation'):
        fast.D(Conditions(1723.15))


def test_fixed_temperature_law_rejects_thermal_mc_before_drawing():
    from diffusor.fitting.montecarlo import run,UncertaintyBudget
    model=DiffusionModel(get('mt_Zr_sievwright2020'),Conditions(1423.15,log_fo2_bar=-8),
            InitialCondition('step',{'x0':0,'C_left':1,'C_right':0}))
    with pytest.raises(ValueError,match='one temperature'):
        run(model,[-1,0,1],[1,.5,0],budget=UncertaintyBudget(sigma_T_K=10),n_draws=2)


@pytest.mark.parametrize('geometry',[1,2])
def test_explicit_radial_centre_respects_positivity_and_conservation(geometry):
    from diffusor.solvers.numerical import solve_1d
    x=np.arange(10.)
    c=np.zeros(10);c[0]=1
    r=solve_1d(x,c,lambda C,T:1.,.5,m=geometry,T_K=1000,theta_time=0.,courant=.5)
    assert np.min(r.C_final)>=0
    assert r.mass_final==pytest.approx(r.mass_initial,rel=1e-12)
    at_zero=solve_1d(x,c,lambda C,T:1.,0,m=geometry,T_K=1000)
    assert at_zero.mass_final==at_zero.mass_initial


def test_history_and_array_initial_conditions_survive_export():
    from diffusor.fitting.fit import fit_time
    from diffusor.dataio.export import result_dict,methods_paragraph
    m=DiffusionModel(get('qz_Ti_cherniak2007'),Conditions(1100,axis='c'),
                    InitialCondition('step',{'x0':0,'C_left':1,'C_right':0}),
                    history=ThermalHistory.isothermal(1373.15,100))
    x=np.linspace(-80,80,41)
    r=fit_time(m,x,m.profile(1e8,x),t_min=1e7,t_max=1e9)
    meta=result_dict(r)
    assert meta['model']['thermal_history']['temperatures_K']==[1373.15,1373.15]
    assert 'overrides the nominal temperature' in methods_paragraph(r)


def test_disabled_composition_dependence_is_consistent_with_isothermal_history():
    c=get('ol_FeMg_dohmen_chakraborty2007_tamed')
    m=DiffusionModel(c,Conditions(1423.15,log_fo2_bar=-12,X={'XFe':.1},axis='c'),
                    InitialCondition('step',{'x0':0,'C_left':.2,'C_right':.3}),
                    comp_key='XFe',composition_dependent=False)
    expected=m._effective_Dt(10,C_ref=.25)
    m.history=ThermalHistory.isothermal(1423.15,10)
    assert m._effective_Dt(10,C_ref=.25)==pytest.approx(expected)
